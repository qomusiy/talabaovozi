"""Import students / teachers / teaching assignments / a timetable from Excel (uploaded in the dashboard).

Column headers are matched flexibly, so real exports (HEMIS etc.) can be used as they are, as long as their
headers resemble one of the aliases below. If a real file uses a header that isn't recognised, the error lists
the columns found; add the header to the alias list - no other code needs to change.
Every kind has a template (app.importers.templates) with headers that are guaranteed to work.
"""
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.faculty import display_code, group_key
from app.importers.templates import INFO_SHEETS
from app.models import Department, Group, ImportLog, Semester, Student, Subject, Teacher, TeachingAssignment
from app.security import (hash_passport, is_valid_passport, normalize_hemis, normalize_passport,
                          temporary_passport)
from app.services.teaching import name_key, teacher_name_keys
from app.timetable.edupage import split_subject


def _norm(h) -> str:
    """'O‘qituvchi F.I.Sh.' -> 'oqituvchifish'"""
    return re.sub(r"[^0-9a-zа-яёқғўҳ]", "", str(h or "").lower())


# field -> aliases (compared after _norm). Exact matches win; then "header contains alias" for aliases >= 4 chars.
STUDENT_COLUMNS = {
    "hemis_id": ["hemisid", "hemis", "talabaid", "studentid", "idraqami", "talabaraqami", "id"],
    "full_name": ["fish", "fio", "fullname", "familiyaismisharifi", "ism", "talabafish"],
    "passport": ["pasport", "passport", "pasportseriyavaraqami", "паспорт", "jshshir"],
    "group": ["guruh", "group", "группа", "guruhnomi"],
    "course": ["kurs", "course", "курс", "bosqich", "year"],
    "specialty": ["mutaxassislik", "yonalish", "specialty", "специальность", "направление"],
    "language": ["talimtili", "til", "language", "языкобучения"],
}
TEACHER_COLUMNS = {
    "external_id": ["oqituvchiid", "teacherid", "id", "xodimid"],
    "full_name": ["fish", "fio", "fullname", "familiyaismisharifi", "oqituvchi"],
    "short_name": ["jadvaldagiism", "qisqaism", "shortname"],
    "department": ["kafedra", "department", "кафедра"],
    "position": ["lavozim", "position", "должность"],
    "room_number": ["xonaraqami", "xona", "room", "roomnumber", "кабинет", "аудитория"],
}
ASSIGNMENT_COLUMNS = {
    "teacher_id": ["oqituvchiid", "teacherid"],
    "teacher_name": ["oqituvchifish", "oqituvchi", "teacher", "teachername", "fish"],
    "group": ["guruh", "group", "группа"],
    "subject": ["fan", "subject", "fannomi", "предмет", "дисциплина"],
    "lesson_type": ["mashgulotturi", "turi", "lessontype", "type", "видзанятия"],
}
TIMETABLE_COLUMNS = {
    "day": ["kun", "hafta kuni", "day", "день"],
    "period": ["juftlik", "para", "period", "пара"],
    "group": ["guruh", "guruhlar", "group", "группа"],
    "subject": ["fan", "subject", "предмет"],
    "lesson_type": ["turi", "mashgulotturi", "type", "видзанятия"],
    "teacher": ["oqituvchi", "teacher", "преподаватель"],
    "room": ["xona", "auditoriya", "room", "аудитория"],
    "week": ["hafta", "week", "неделя"],
}
LESSON_TYPE_MAP = {
    "maruza": "lecture", "lecture": "lecture", "лекция": "lecture", "m": "lecture", "ma": "lecture",
    "amaliy": "practice", "amaliyot": "practice", "practice": "practice", "практика": "practice", "a": "practice",
    "laboratoriya": "lab", "lab": "lab", "лабораторная": "lab", "l": "lab",
    "seminar": "seminar", "семинар": "seminar", "s": "seminar",
}
LANGUAGE_MAP = {"ozbek": "uz", "uzbek": "uz", "uz": "uz", "rus": "ru", "russian": "ru", "ru": "ru",
                "русский": "ru", "ingliz": "en", "english": "en", "en": "en", "ing": "en",
                "turkman": "tk", "tk": "tk"}
DAY_MAP = {**{d: i for i, d in enumerate(["dushanba", "seshanba", "chorshanba", "payshanba", "juma", "shanba"])},
           **{d: i for i, d in enumerate(["monday", "tuesday", "wednesday", "thursday", "friday", "saturday"])},
           **{d: i for i, d in enumerate(["понедельник", "вторник", "среда", "четверг", "пятница", "суббота"])},
           **{str(i + 1): i for i in range(6)}}
# Used when an uploaded timetable has no times (TIIAME's periods).
DEFAULT_PERIODS = [{"n": 1, "start": "08:00", "end": "09:20"}, {"n": 2, "start": "09:30", "end": "10:50"},
                   {"n": 3, "start": "11:00", "end": "12:20"}, {"n": 4, "start": "13:00", "end": "14:20"},
                   {"n": 5, "start": "14:30", "end": "15:50"}, {"n": 6, "start": "16:00", "end": "17:20"}]
DAY_NAMES = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba"]


# Problems are codes + values; the dashboard shows them in its language (imp_* / impf_* texts), the CLI in English.
ROW_MESSAGES = {
    "bad_hemis": "invalid HEMIS id '{value}'",
    "dup_hemis": "HEMIS id {value} appears twice in the file",
    "new_needs": "new student {value} needs a name and a group",
    "bad_passport": "invalid passport code '{value}' (expected like AD1231212)",
    "no_name": "empty name",
    "no_teacher": "no teacher",
    "bad_slot": "day '{day}' / period '{period}' not understood",
    "no_rows": "the file has no rows",
}
FILE_MESSAGES = {
    "open": "Could not open the file as .xlsx.",
    "empty": "The file is empty.",
    "too_big": "The file is larger than 10 MB.",
    "missing": "Missing column(s): {missing}. Columns in the file: {found}",
}
# Column names as in the templates, used in "missing column" messages.
FIELD_LABELS = {"hemis_id": "HEMIS ID", "full_name": "F.I.Sh.", "passport": "Pasport", "group": "Guruh",
                "subject": "Fan", "day": "Kun", "period": "Juftlik"}


class ImportFileError(Exception):
    """The file as a whole can't be imported (wrong format, missing columns)."""

    def __init__(self, code: str, **params):
        self.code, self.params = code, params
        super().__init__(FILE_MESSAGES[code].format(**params))


@dataclass
class ImportReport:
    kind: str
    created: int = 0
    updated: int = 0
    skipped: int = 0
    temporary: int = 0     # students given a temporary passport code
    deactivated: int = 0   # students not in the file (with deactivate_missing)
    errors: list[tuple[int, str, dict]] = field(default_factory=list)  # (row, code, values)

    def error(self, row: int, code: str, **values) -> None:
        self.errors.append((row, code, values))

    def messages(self) -> list[str]:
        return [f"Row {row}: {ROW_MESSAGES[code].format(**values)}" for row, code, values in self.errors]

    def summary(self) -> str:
        parts = [f"{self.created} created", f"{self.updated} updated"]
        if self.skipped:
            parts.append(f"{self.skipped} duplicates skipped")
        if self.temporary:
            parts.append(f"{self.temporary} with a temporary passport code")
        if self.deactivated:
            parts.append(f"{self.deactivated} deactivated")
        parts.append(f"{len(self.errors)} errors")
        return f"{self.kind}: " + ", ".join(parts)

    def log(self, filename: str, admin: str | None) -> ImportLog:
        note = self.summary() + "".join(f"\n{e}" for e in self.messages()[:20])
        return ImportLog(kind=self.kind, filename=filename[:255], created=self.created, updated=self.updated,
                         errors=len(self.errors), note=note, admin=admin)


def _map_columns(headers: list, spec: dict[str, list[str]], required: list[str]) -> dict[str, int]:
    normed = [_norm(h) for h in headers]
    mapping: dict[str, int] = {}
    taken: set[int] = set()
    # Pass 1: exact matches. Pass 2: substring matches for longer aliases.
    for exact in (True, False):
        for fld, aliases in spec.items():
            if fld in mapping:
                continue
            for alias in (_norm(a) for a in aliases):
                hit = next((i for i, h in enumerate(normed) if i not in taken and h and (
                    h == alias if exact else (len(alias) >= 4 and alias in h))), None)
                if hit is not None:
                    mapping[fld] = hit
                    taken.add(hit)
                    break
    missing = [f for f in required if f not in mapping]
    if missing:
        raise ImportFileError("missing", missing=", ".join(FIELD_LABELS.get(f, f) for f in missing),
                              found=", ".join(str(h) for h in headers if h) or "—")
    return mapping


def _sheets(source: Path | BinaryIO) -> list[list[tuple]]:
    """The rows of every sheet to read: the active one first, then the others - but never the template's
    instructions, example and drop-down sheets (app.importers.templates.INFO_SHEETS)."""
    if hasattr(source, "seek"):
        source.seek(0)
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError("open") from e
    sheets = [wb.active, *(ws for ws in wb.worksheets if ws is not wb.active)]
    out = [list(ws.iter_rows(values_only=True)) for ws in sheets if ws.title not in INFO_SHEETS]
    wb.close()
    return out


def _rows_of(rows: list[tuple], spec, required) -> list[tuple[int, dict]]:
    """[(excel_row_number, {field: value})] of one sheet. Finds the header row within the first 10 rows;
    raises ImportFileError("missing") if the sheet doesn't have the needed columns."""
    last_err: ImportFileError | None = None
    for header_idx, header in enumerate(rows[:10]):
        try:
            mapping = _map_columns(list(header), spec, required)
        except ImportFileError as e:
            last_err = last_err or e
            continue
        out = []
        for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            values = {f: (row[i] if i < len(row) else None) for f, i in mapping.items()}
            values = {k: (str(v).strip() if v is not None and str(v).strip() != "" else None)
                      for k, v in values.items()}
            if any(values.values()):
                out.append((offset, values))
        return out
    raise last_err or ImportFileError("empty")


def _read_rows(source: Path | BinaryIO, spec, required) -> list[tuple[int, dict]]:
    """The rows of the first sheet that has the needed columns (the active sheet first)."""
    last_err: ImportFileError | None = None
    for rows in _sheets(source):
        try:
            return _rows_of(rows, spec, required)
        except ImportFileError as e:
            last_err = last_err or e
    raise last_err or ImportFileError("empty")


def teaching_sheets(source) -> tuple[list, list]:
    """(teacher rows, assignment rows) of a file with either or both: a sheet with Guruh + Fan columns is
    assignments, one with only a name column is teachers (the combined template has both)."""
    teachers: list = []
    assignments: list = []
    found = False
    for rows in _sheets(source):
        try:
            assignments += _rows_of(rows, ASSIGNMENT_COLUMNS, ["group", "subject"])
            found = True
            continue
        except ImportFileError:
            pass
        try:
            teachers += _rows_of(rows, TEACHER_COLUMNS, ["full_name"])
            found = True
        except ImportFileError:
            pass
    if not found:
        raise ImportFileError("missing", missing="F.I.Sh. / O‘qituvchi, Guruh, Fan", found="—")
    return teachers, assignments


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def tidy_name(name: str) -> str:
    """'TOSHMATOV ESHMAT O‘G‘LI' -> 'Toshmatov Eshmat O‘g‘li' (names typed in capitals only)."""
    name = _clean(name)
    if name != name.upper():
        return name
    return " ".join(w[:1] + w[1:].lower() for w in name.split())


def _course(raw: str | None, group_name: str | None) -> int | None:
    """'2-kurs' -> 2; otherwise from the admission year in the group name ('25-16 IQT' -> 2nd year in 2026/27)."""
    if raw and (m := re.search(r"\d", raw)):
        return int(m.group())
    m = re.match(r"^(\d{2})-", group_name or "")
    if not m:
        return None
    today = date.today()
    school_year = today.year if today.month >= 8 else today.year - 1
    course = school_year - (2000 + int(m.group(1))) + 1
    return course if 1 <= course <= 6 else None


class _Groups:
    """Finds groups by their matching key (see app.faculty.group_key), creating missing ones."""

    def __init__(self, session: AsyncSession):
        self.session, self.by_key, self.by_name = session, None, None

    async def get(self, name: str, **attrs) -> Group:
        if self.by_key is None:
            rows = list((await self.session.execute(select(Group))).scalars())
            self.by_key = {g.key or group_key(g.name) or g.name: g for g in rows}
            self.by_name = {g.name: g for g in rows}
        name = _clean(name)
        key = group_key(name) or name
        group = self.by_key.get(key) or self.by_name.get(name)
        if group is None:
            group = Group(name=name, key=group_key(name), code=display_code(name), course=_course(None, name))
            self.session.add(group)
            self.by_key[key] = self.by_name[name] = group
        for k, v in attrs.items():
            if v is not None and (k != "course" or not group.course):
                setattr(group, k, v)
        return group


async def import_students(session: AsyncSession, source, filename: str = "", deactivate_missing: bool = False,
                          admin: str | None = None) -> ImportReport:
    """Upsert students by HEMIS id. Existing Telegram links are kept.

    Only the HEMIS column is required: a file with just HEMIS + passport updates the passport codes.
    New students need a name and a group. Students without a passport code get a temporary one
    (app.security.temporary_passport) until a file with the real codes is uploaded.
    deactivate_missing=True marks students not in the file as inactive (graduated, expelled...).
    """
    report = ImportReport("students")
    rows = _read_rows(source, STUDENT_COLUMNS, ["hemis_id"])
    existing = {s.hemis_id: s for s in (await session.execute(select(Student))).scalars()}
    groups = _Groups(session)
    seen: set[str] = set()

    for row_no, r in rows:
        hemis = normalize_hemis(r["hemis_id"] or "")
        if not re.fullmatch(r"\d{5,20}", hemis):
            report.error(row_no, "bad_hemis", value=r["hemis_id"])
            continue
        if hemis in seen:
            report.error(row_no, "dup_hemis", value=hemis)
            continue
        student = existing.get(hemis)
        if student is None and not (r.get("full_name") and r.get("group")):
            report.error(row_no, "new_needs", value=hemis)
            continue
        passport = normalize_passport(r["passport"]) if r.get("passport") else None
        if passport and not is_valid_passport(passport):
            report.error(row_no, "bad_passport", value=r["passport"])
            continue
        seen.add(hemis)

        if student is None:
            student = Student(hemis_id=hemis)
            session.add(student)
            existing[hemis] = student
            report.created += 1
        else:
            report.updated += 1
        if passport:
            student.passport_hash, student.passport_temporary = hash_passport(passport), False
        elif not student.passport_hash or student.passport_temporary:
            student.passport_hash, student.passport_temporary = hash_passport(temporary_passport(hemis)), True
            report.temporary += 1
        if r.get("full_name"):
            student.full_name = tidy_name(r["full_name"])
        if r.get("group"):
            student.group = await groups.get(r["group"], course=_course(r.get("course"), r["group"]),
                                             specialty=r.get("specialty"),
                                             language=LANGUAGE_MAP.get(_norm(r.get("language") or "")))
        student.is_active = True

    if deactivate_missing and not report.errors:
        for hemis, student in existing.items():
            if hemis not in seen and student.is_active:
                student.is_active = False
                report.deactivated += 1

    session.add(report.log(filename, admin))
    await session.flush()
    return report


class _Teachers:
    """Finds teachers by Excel id, full name, or any way a timetable writes their name ('Qodirov RS')."""

    def __init__(self, teachers: list[Teacher]):
        self.by_ext = {t.external_id: t for t in teachers if t.external_id}
        self.by_name: dict[str, Teacher] = {}
        for t in teachers:
            self.add(t)

    def add(self, t: Teacher) -> None:
        for k in teacher_name_keys(t.full_name) | ({name_key(t.short_name)} if t.short_name else set()):
            self.by_name.setdefault(k, t)
        if t.external_id:
            self.by_ext[t.external_id] = t

    def find(self, ext: str | None = None, name: str | None = None, short: str | None = None) -> Teacher | None:
        if ext and (t := self.by_ext.get(ext)):
            return t
        for candidate in (name, short):
            if candidate and (t := self.by_name.get(name_key(candidate))):
                return t
        return None


async def import_teachers(session: AsyncSession, source, filename: str = "", admin: str | None = None,
                          rows: list | None = None) -> ImportReport:
    """Upsert teachers. Matches the teachers taken from the timetable, so this adds full names, departments,
    positions and rooms to them. Rows without a match are created. `rows`: already read (teaching_sheets)."""
    report = ImportReport("teachers")
    rows = _read_rows(source, TEACHER_COLUMNS, ["full_name"]) if rows is None else rows
    teachers = _Teachers(list((await session.execute(select(Teacher))).scalars()))
    departments = {d.name: d for d in (await session.execute(select(Department))).scalars()}

    for row_no, r in rows:
        name = tidy_name(r["full_name"] or "")
        if not name:
            report.error(row_no, "no_name")
            continue
        ext = normalize_hemis(r["external_id"]) if r.get("external_id") else None
        teacher = teachers.find(ext, name, r.get("short_name"))
        if teacher is None:
            teacher = Teacher(full_name=name, short_name=r.get("short_name"))
            session.add(teacher)
            report.created += 1
        else:
            report.updated += 1
        teacher.full_name = name
        teacher.external_id = ext or teacher.external_id
        teacher.short_name = r.get("short_name") or teacher.short_name
        teacher.position = r.get("position") or teacher.position
        teacher.room_number = r.get("room_number") or teacher.room_number
        teacher.is_active = True
        if dept_name := r.get("department"):
            dept_name = _clean(dept_name)
            if dept_name not in departments:
                departments[dept_name] = Department(name=dept_name)
                session.add(departments[dept_name])
            teacher.department = departments[dept_name]
        teachers.add(teacher)

    session.add(report.log(filename, admin))
    await session.flush()
    return report


async def import_assignments(session: AsyncSession, source, semester: Semester, filename: str = "",
                             admin: str | None = None, rows: list | None = None) -> ImportReport:
    """Replace the semester's teaching assignments with the contents of the file.

    Teachers are matched by id, full name or timetable name; unknown teachers are created.
    Existing feedback is not affected (it references teachers, not assignments).
    """
    report = ImportReport("assignments")
    rows = _read_rows(source, ASSIGNMENT_COLUMNS, ["group", "subject"]) if rows is None else rows
    teachers = _Teachers(list((await session.execute(select(Teacher))).scalars()))
    subjects = {s.name: s for s in (await session.execute(select(Subject))).scalars()}
    groups = _Groups(session)
    seen: set[tuple] = set()
    new_rows: list[TeachingAssignment] = []

    for row_no, r in rows:
        ext = normalize_hemis(r["teacher_id"]) if r.get("teacher_id") else None
        teacher = teachers.find(ext, r.get("teacher_name"))
        if teacher is None:
            if not r.get("teacher_name"):
                report.error(row_no, "no_teacher")
                continue
            teacher = Teacher(full_name=tidy_name(r["teacher_name"]), short_name=_clean(r["teacher_name"]),
                              external_id=ext)
            session.add(teacher)
            teachers.add(teacher)
        subject_name, suffix_type = split_subject(_clean(r["subject"]))
        if subject_name not in subjects:
            subjects[subject_name] = Subject(name=subject_name)
            session.add(subjects[subject_name])
        lesson_type = LESSON_TYPE_MAP.get(_norm(r.get("lesson_type") or ""), suffix_type)
        group = await groups.get(r["group"])

        key = (id(teacher), group.name, subject_name, lesson_type)
        if key in seen:
            report.skipped += 1
            continue
        seen.add(key)
        new_rows.append(TeachingAssignment(semester=semester, teacher=teacher, group=group,
                                           subject=subjects[subject_name], lesson_type=lesson_type))

    if not new_rows:
        if not report.errors:
            report.error(0, "no_rows")
        return report
    await session.execute(delete(TeachingAssignment).where(TeachingAssignment.semester_id == semester.id))
    session.add_all(new_rows)
    report.created = len(new_rows)
    session.add(report.log(filename, admin))
    await session.flush()
    return report


def parse_timetable(source) -> tuple[dict, ImportReport]:
    """An uploaded timetable -> the same structure as an EduPage download (see app.timetable.edupage.parse).

    One row per class: day, period, group(s), subject, type, teacher(s), room, week (A/B, empty = every week).
    Several groups or teachers in one cell are separated by commas. Every group in the file counts as the faculty's.
    """
    report = ImportReport("timetable")
    rows = _read_rows(source, TIMETABLE_COLUMNS, ["day", "period", "group", "subject"])
    groups: dict[str, dict] = {}
    teachers: dict[str, dict] = {}
    lessons = []
    for row_no, r in rows:
        day = DAY_MAP.get(_norm(r["day"]).replace("kuni", "")) if r.get("day") else None
        period = int(m.group()) if r.get("period") and (m := re.search(r"\d+", r["period"])) else None
        if day is None or not period or period > len(DEFAULT_PERIODS):
            report.error(row_no, "bad_slot", day=r.get("day") or "", period=r.get("period") or "")
            continue
        gids = []
        for name in filter(None, (_clean(x) for x in r["group"].split(","))):
            gid = (group_key(name) or name).replace(" ", "-")
            groups.setdefault(gid, {"id": gid, "name": name, "course": _course(None, name),
                                    "specialty": display_code(name) or "", "lang": "uz", "note": "", "faculty": True})
            gids.append(gid)
        tids = []
        for name in filter(None, (_clean(x) for x in (r.get("teacher") or "").split(","))):
            tid = name_key(name).replace(" ", "-")
            teachers.setdefault(tid, {"id": tid, "name": name, "faculty": True})
            tids.append(tid)
        subject, suffix_type = split_subject(_clean(r["subject"]))
        week = (r.get("week") or "").strip().upper()[:1]
        lessons.append({"day": day, "period": period, "week": week if week in ("A", "B") else "",
                        "subject": subject, "type": LESSON_TYPE_MAP.get(_norm(r.get("lesson_type") or ""), suffix_type),
                        "room": _clean(r.get("room") or ""), "teachers": tids, "groups": gids})
    report.created = len(lessons)
    data = {"source": "excel", "periods": DEFAULT_PERIODS, "days": DAY_NAMES, "groups": list(groups.values()),
            "teachers": sorted(teachers.values(), key=lambda t: t["name"].lower()),
            "lessons": sorted(lessons, key=lambda x: (x["day"], x["period"], x["week"])), "unscheduled": 0}
    return data, report
