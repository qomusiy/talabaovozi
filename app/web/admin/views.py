"""Admin dashboard pages."""
import logging
import math
import secrets
import time
from datetime import datetime
from io import BytesIO
from typing import Annotated
from urllib.parse import quote, urlencode
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BeforeValidator
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.faculty import specialty_name
from app.models import RATING_CRITERIA, Admin, Teacher
from app.security import temporary_passport, verify_password
from app.services.teaching import refresh_from_edupage
from app.timetable import view as ttv
from app.web.admin import analytics as an
from app.web.admin import display
from app.web.admin.i18n import DAY_NAMES, DEFAULT_LANG, LANGS, js_texts, translator
from app.web.deps import current_admin, get_session, superadmin
from app.web.security import client_ip, throttle
from app.web.templating import templates

router = APIRouter(prefix="/admin")
log = logging.getLogger("admin")
LANG_COOKIE = "admin_lang"
STATUSES = ("not_registered", "registered", "partial", "completed")
TASHKENT = ZoneInfo("Asia/Tashkent")


# Filter dropdowns send "" for "all" (?group=&teacher=): a blank or non-numeric filter means no filter.
OptionalId = Annotated[int | None, BeforeValidator(lambda v: int(v) if str(v).strip().isdigit() else None)]
PageNumber = Annotated[int, BeforeValidator(lambda v: max(1, int(v)) if str(v).strip().isdigit() else 1)]


def get_lang(request: Request) -> str:
    lang = request.cookies.get(LANG_COOKIE, DEFAULT_LANG)
    return lang if lang in LANGS else DEFAULT_LANG


def criteria_labels(t) -> dict[str, str]:
    return {c: t(f"c_{c}") for c in RATING_CRITERIA}


def delta(current, previous) -> float | None:
    if current is None or previous is None:
        return None
    return round(current - previous, 2)


async def resolve_campaign(session: AsyncSession, raw: str | None) -> tuple[int | None, list]:
    """'all' -> None (all campaigns); missing -> the most recent campaign."""
    campaigns = await an.campaigns(session)
    if raw == "all":
        return None, campaigns
    if raw and raw.isdigit() and any(c.id == int(raw) for c in campaigns):
        return int(raw), campaigns
    return (campaigns[0].id if campaigns else None), campaigns


# ---------------------------------------------------------------- Breadcrumbs
# The browser remembers the path taken to a page (static/crumbs.js) and shows it as "Teachers › Name › Timetable".
# Each page describes itself; labels carry their text key so they can be shown again later in another language.
SECTIONS = {"/admin": "nav_overview", "/admin/teachers": "nav_teachers", "/admin/groups": "nav_groups",
            "/admin/students": "nav_students", "/admin/comments": "nav_comments", "/admin/timetable": "tt_title",
            "/admin/surveys": "nav_surveys", "/admin/data": "nav_data", "/admin/records": "nav_records", "/admin/audit": "nav_audit",
            "/admin/bot": "nav_bot", "/admin/admins": "nav_admins", "/admin/account": "nav_account",
            "/admin/chats": "nav_chats"}
CRUMB_TEXTS = (*SECTIONS.values(), "tt_feedback", "students_of")


def label(t, key: str, **params) -> dict:
    return {"s": t(key, **params), "k": key, "v": params}


def plain(text: str) -> dict:
    return {"s": text}


def crumb(label_: dict, kind: str, entity: str | None = None, ctx: dict | None = None) -> dict:
    """This page in the breadcrumbs. `entity` names what it shows ("teacher:12"); when the previous step shows the
    same thing, `ctx` is used instead of `label` (a teacher's timetable opened from their page reads "Timetable")."""
    return {"label": label_, "kind": kind, "entity": entity, "ctx": ctx}


def campaign_url(campaign_id: int | None):
    def url(path: str, **params) -> str:
        """Link that keeps the selected campaign."""
        q = {"campaign": campaign_id or "all", **{k: v for k, v in params.items() if v not in (None, "")}}
        return f"{path}?{urlencode(q)}"
    return url


def section(t, url, path: str, **params) -> dict:
    """A top-level page as a parent step, for pages opened directly (bookmark, new tab)."""
    return {"label": label(t, SECTIONS[path]), "kind": path, "path": path, "url": url(path, **params)}


async def render(request: Request, template: str, admin: Admin, campaign: tuple[int | None, list],
                 crumb: dict | None = None, parents: list[dict] | None = None, **ctx) -> HTMLResponse:
    campaign_id, campaigns = campaign
    selected = next((c for c in campaigns if c.id == campaign_id), None)
    lang = get_lang(request)
    t = translator(lang)
    url = campaign_url(campaign_id)

    def num(n) -> str:
        """1,234 in English; 1 234 in Uzbek/Russian."""
        s = f"{n:,}"
        return s if lang == "en" else s.replace(",", " ")

    path = request.url.path
    current_url = path + (f"?{request.url.query}" if request.url.query else "")
    me = {"kind": path, "entity": None, "ctx": None, **(crumb or {"label": label(t, SECTIONS.get(path, "nav_overview"))}),
          "path": path, "url": current_url}
    crumbs = {"self": me, "fallback": [*(parents or []), me], "texts": {k: t(k) for k in CRUMB_TEXTS}}
    request.state.audit_title = " › ".join(c["label"]["s"] for c in crumbs["fallback"])[:255]  # for the activity log
    return templates.TemplateResponse(request, template, {
        "admin": admin, "campaigns": campaigns, "campaign_id": campaign_id, "campaign": selected,
        "criteria": criteria_labels(t), "url": url, "path": path, "t": t, "lang": lang, "crumbs": crumbs,
        "specialty": lambda code: specialty_name(code, lang),
        "langs": LANGS, "num": num, "js_texts": js_texts(lang), "current_url": quote(current_url, safe=""),
        "is_super": admin.is_super, "is_programmer": admin.is_programmer, "show_who": admin.sees_identity,
        "flash": request.session.pop("flash", None), "now_hm": datetime.now(TASHKENT).strftime("%H:%M"),
        "display": display.clean(admin.prefs), "display_attrs": display.html_attrs(admin.prefs),
        "demo_active": any(c.is_demo for c in campaigns), **ctx,
    })


def without_students(rows: list[dict], t) -> list[dict]:
    """For viewers who may not see who rated (Adminlar › "show who rated" off): the student's name becomes
    "Talaba"; the group stays. Done here, so the names are not in the page at all (not even its data)."""
    return [{**r, "student": t("anon_student"), "hemis": None} for r in rows]


def department_groups(departments: list, t) -> list[tuple[str, list]]:
    """Kafedras by faculty, for <optgroup>s: this faculty's first, then the others (Oliy matematika, ...)."""
    home = get_settings().edupage_faculty
    by: dict[str, list] = {}
    for d in departments:
        by.setdefault(d.faculty or home, []).append(d)
    return [(t("faculty_of", name=f), by[f]) for f in sorted(by, key=lambda f: (f != home, f.lower()))]


def flash(request: Request, kind: str, text: str) -> None:
    """A message shown once at the top of the next page (kind: good / bad / info)."""
    request.session["flash"] = {"kind": kind, "text": text}


# ---------------------------------------------------------------- Language & auth

@router.get("/lang/{code}")
async def set_language(code: str, next: str = "/admin"):
    target = next if next.startswith("/admin") else "/admin"
    response = RedirectResponse(target, status_code=303)
    if code in LANGS:
        response.set_cookie(LANG_COOKIE, code, max_age=365 * 24 * 3600, samesite="lax")
    return response


def login_context(request: Request, next: str, error: bool = False, username: str | None = None) -> dict:
    lang = get_lang(request)
    t = translator(lang)
    return {"next": next, "error": t("login_error") if error else None, "username": username,
            "year": datetime.now().year, "t": t, "lang": lang, "langs": LANGS,
            "help": t("login_help", contact=get_settings().tech_contact),
            "current_url": quote(f"/admin/login?next={next}", safe="")}


@router.get("/login", response_class=HTMLResponse)
async def login_page(request: Request, next: str = "/admin"):
    return templates.TemplateResponse(request, "admin/login.html", login_context(request, next))


@router.post("/login", response_class=HTMLResponse)
async def login(request: Request, username: str = Form(...), password: str = Form(...), next: str = Form("/admin"),
                session: AsyncSession = Depends(get_session, scope="function")):
    keys = (f"ip:{client_ip(request)}", f"user:{username.strip().lower()}")
    if wait := throttle.blocked_for(*keys):
        ctx = login_context(request, next, username=username)
        ctx["error"] = translator(get_lang(request))("login_blocked", n=wait)
        return templates.TemplateResponse(request, "admin/login.html", ctx, status_code=429)
    admin = (await session.execute(select(Admin).where(Admin.username == username.strip()))).scalar_one_or_none()
    if admin is None or not admin.is_active or not verify_password(password, admin.password_hash):
        throttle.failed(*keys)
        return templates.TemplateResponse(request, "admin/login.html",
                                          login_context(request, next, error=True, username=username),
                                          status_code=401)
    throttle.succeeded(*keys)
    request.session["admin_id"] = admin.id
    request.session["sid"] = secrets.token_hex(8)       # ties this login's activity together in the log
    request.session["login_at"] = int(time.time())
    return RedirectResponse(next if next.startswith("/admin") else "/admin", status_code=303)


@router.get("/logout")
async def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/admin/login", status_code=303)


# ---------------------------------------------------------------- Pages

def greeting_key() -> str:
    hour = datetime.now(TASHKENT).hour
    return "greet_morning" if 5 <= hour < 12 else "greet_day" if hour < 17 else "greet_evening" if hour < 23 else "greet_night"


@router.get("", response_class=HTMLResponse)
async def overview(request: Request, campaign: str | None = None, admin: Admin = Depends(current_admin),
                   session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    lang = get_lang(request)
    t = translator(lang)
    teachers = await an.teacher_table(session, cid)
    groups = await an.group_table(session, cid)
    kpis = await an.kpis(session, cid)
    completed = sum(g["completed"] for g in groups)
    kpis |= {"completed": completed,
             "completed_pct": round(100 * completed / kpis["students"], 1) if kpis["students"] else 0,
             "comments_pct": round(100 * kpis["comments"] / kpis["evaluations"], 1) if kpis["evaluations"] else 0}
    prev_id = await an.previous_campaign_id(session, cid)
    prev = await an.kpis(session, prev_id) if prev_id else None
    deltas = {k: delta(kpis[k], prev[k]) for k in ("overall", *RATING_CRITERIA)} if prev else {}

    selected = next((c for c in campaigns if c.id == cid), None)
    status = None
    if selected:
        now = datetime.now(TASHKENT)
        closes = selected.closes_at.astimezone(TASHKENT)
        status = {"open": selected.opens_at <= now < closes, "closes": closes.strftime("%d.%m.%Y"),
                  "days_left": max(0, (closes.date() - now.date()).days)}

    # Ranking: top / bottom 10 by score (at least 5 evaluations, so one angry student can't decide it), most rated.
    rated = [x for x in teachers if x["n"] >= 5 and x["overall"] is not None]
    most = sorted((x for x in teachers if x["n"]), key=lambda x: -x["n"])[:10]
    ranking = {"best": sorted(rated, key=lambda x: -x["overall"])[:10],
               "worst": sorted(rated, key=lambda x: x["overall"])[:10], "most": most}
    return await render(
        request, "admin/overview.html", admin, (cid, campaigns),
        greeting=t(greeting_key()), kpis=kpis, deltas=deltas, prev=prev, status=status,
        ranking=ranking, most_n=most[0]["n"] if most else 1,
        attention=[x for x in reversed(rated) if x["overall"] < 6][:6],
        comments=(lambda rows: rows if admin.sees_identity else without_students(rows, t))(
            await an.latest_comments(session, cid, limit=5)),
        data={
            "daily": await an.daily_submissions(session, cid),
            "curve": await an.participation_curve(session, cid),
            "students": kpis["students"],
            "specialties": [{**x, "name": specialty_name(x["code"], lang)}
                            for x in await an.specialty_criteria(session, cid)],
            "distribution": await an.distribution(session, cid),
            "criteria": criteria_labels(t),
        },
    )


@router.get("/teachers", response_class=HTMLResponse)
async def teachers(request: Request, campaign: str | None = None, department: OptionalId = None,
                   admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    t = translator(get_lang(request))
    return await render(request, "admin/teachers.html", admin, (cid, campaigns),
                        rows=await an.teacher_table(session, cid, department),
                        departments=department_groups(await an.departments(session), t), department_id=department)


@router.get("/teachers/{teacher_id}", response_class=HTMLResponse)
async def teacher(request: Request, teacher_id: int, campaign: str | None = None,
                  admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    detail = await an.teacher_detail(session, teacher_id, cid)
    if detail is None:
        raise HTTPException(404)
    t = translator(get_lang(request))
    if not admin.sees_identity:
        detail["feedback"] = without_students(detail["feedback"], t)
    tt = await ttv.latest(session)
    return await render(request, "admin/teacher.html", admin, (cid, campaigns),
                        crumb(plain(detail["teacher"]["name"]), "teacher", f"teacher:{teacher_id}", label(t, "tt_feedback")),
                        [section(t, campaign_url(cid), "/admin/teachers")],
                        d=detail, data={**detail, "criteria": criteria_labels(t)},
                        timetable_id=tt_teacher_id(tt, detail["teacher"]) if tt else None)


@router.get("/groups", response_class=HTMLResponse)
async def groups(request: Request, campaign: str | None = None,
                 admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    """Charts and figures cover every group; the course filter and search only narrow the table (admin.js)."""
    cid, campaigns = await resolve_campaign(session, campaign)
    rows = await an.group_table(session, cid)
    students = sum(g["students"] for g in rows)
    active = [g for g in rows if g["students"]]
    summary = {
        "groups": len(rows),
        "avg": round(100 * sum(g["participants"] for g in rows) / students, 1) if students else 0,
        "best": max(active, key=lambda g: g["participation_pct"], default=None),
        "worst": min(active, key=lambda g: g["participation_pct"], default=None),
    }
    return await render(request, "admin/groups.html", admin, (cid, campaigns), rows=rows,
                        courses=sorted({g["course"] for g in rows if g["course"]}), summary=summary, data={"groups": rows, "courses": an.course_summary(rows)})


@router.get("/students", response_class=HTMLResponse)
async def students(request: Request, campaign: str | None = None, group: OptionalId = None,
                   status: str | None = None, q: str | None = None, page: PageNumber = 1,
                   admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    per_page = 50
    result = await an.student_table(session, cid, group, status, q, max(page, 1), per_page)
    groups = await an.all_groups(session)
    chosen = next((g for g in groups if g.id == group), None)
    t = translator(get_lang(request))
    return await render(request, "admin/students.html", admin, (cid, campaigns),
                        crumb(label(t, "students_of", group=chosen.name), "/admin/students", f"group:{chosen.id}")
                        if chosen else None, [section(t, campaign_url(cid), "/admin/groups")],
                        temporary_passport=temporary_passport,
                        result=result, groups=groups, group_id=group, status=status, q=q or "",
                        page=max(page, 1), pages=max(1, math.ceil(result.total / per_page)), statuses=STATUSES)


@router.get("/comments", response_class=HTMLResponse)
async def comments(request: Request, campaign: str | None = None, q: str | None = None, band: str | None = None,
                   teacher: OptionalId = None, group: OptionalId = None, page: PageNumber = 1,
                   admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    per_page = 30
    t = translator(get_lang(request))
    result = await an.comment_feed(session, cid, q, band, teacher, group, max(page, 1), per_page,
                                   by_student=admin.sees_identity)
    if not admin.sees_identity:
        result.rows = without_students(result.rows, t)
    # With a group chosen only its teachers are offered, with a teacher only their groups (others would show nothing).
    links = await an.teacher_group_links(session, cid)
    teachers_ = [x for x in await an.all_teachers(session) if not group or x.id == teacher or (x.id, group) in links]
    groups_ = [g for g in await an.all_groups(session) if not teacher or g.id == group or (teacher, g.id) in links]
    return await render(request, "admin/comments.html", admin, (cid, campaigns), result=result, q=q or "",
                        band_filter=band if band in an.COMMENT_BANDS else None, teacher_id=teacher, group_id=group,
                        teachers=teachers_, groups=groups_,
                        page=max(page, 1), pages=max(1, math.ceil(result.total / per_page)))


# ---------------------------------------------------------------- Timetable (EduPage)

def _name_key(name: str | None) -> str:
    return " ".join((name or "").replace("‘", "'").replace("’", "'").replace("ʻ", "'").lower().split())


def tt_teacher_id(tt: ttv.Timetable, teacher: dict) -> str | None:
    """The timetable id of a teacher from our database (matched by EduPage id or name), if any."""
    names = {_name_key(teacher.get("name")), _name_key(teacher.get("short_name"))} - {""}
    for t in tt.teachers.values():
        if t["faculty"] and (t["id"] == teacher.get("edupage_id") or _name_key(t["name"]) in names):
            return t["id"]
    return None


def timetable_context(request: Request, tt: ttv.Timetable) -> dict:
    lang = get_lang(request)
    fetched = tt.fetched_at.astimezone(TASHKENT).strftime("%d.%m.%Y %H:%M")
    return {"tt": tt, "day_names": DAY_NAMES[lang][:len(tt.days)] if lang in DAY_NAMES else tt.days,
            "updated": fetched}


@router.get("/timetable", response_class=HTMLResponse)
async def timetable(request: Request, campaign: str | None = None, tab: str = "groups",
                    admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    tt = await ttv.latest(session)
    if tt is None:
        return await render(request, "admin/timetable.html", admin, (cid, campaigns), tt=None, tab="groups")
    ctx = timetable_context(request, tt)
    groups = tt.faculty_groups()
    return await render(
        request, "admin/timetable.html", admin, (cid, campaigns), **ctx,
        tab=tab if tab in ("groups", "teachers") else "groups",
        courses=tt.courses(), teachers=tt.faculty_teachers(), mix=tt.mix(tt.faculty_lessons),
        specialties=sorted({(g["specialty"], specialty_name(g["specialty"], get_lang(request))) for g in groups},
                           key=lambda x: x[1].lower()),
        summary={"groups": len(groups), "teachers": sum(x["faculty"] for x in tt.teachers.values()),
                 "pairs": ttv.pairs(tt.faculty_lessons), "subjects": len({x["subject"] for x in tt.faculty_lessons})},
        data={"load": tt.load(), "days": ctx["day_names"], "periods": tt.periods},
    )


@router.get("/timetable/group/{gid}", response_class=HTMLResponse)
async def timetable_group(request: Request, gid: str, campaign: str | None = None,
                          admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    tt = await ttv.latest(session)
    if tt is None or gid not in tt.groups:
        raise HTTPException(404)
    lessons = tt.by_group[gid]
    t = translator(get_lang(request))
    return await render(
        request, "admin/timetable_week.html", admin, (cid, campaigns),
        crumb(plain(tt.groups[gid]["name"]), "tt-group", f"tt-group:{gid}"),
        [section(t, campaign_url(cid), "/admin/timetable", tab="groups")], **timetable_context(request, tt),
        kind="group", entity=tt.groups[gid], week=tt.week(lessons, datetime.now(TASHKENT)), stats=tt.stats(lessons),
        people=len({p for x in lessons for p in x["teachers"]}), subjects=tt.subjects(lessons, "teachers"),
        courses=tt.courses())


@router.get("/timetable/teacher/{tid}", response_class=HTMLResponse)
async def timetable_teacher(request: Request, tid: str, campaign: str | None = None,
                            admin: Admin = Depends(current_admin), session: AsyncSession = Depends(get_session, scope="function")):
    cid, campaigns = await resolve_campaign(session, campaign)
    tt = await ttv.latest(session)
    if tt is None or tid not in tt.teachers:
        raise HTTPException(404)
    teacher = tt.teachers[tid]
    lessons = tt.by_teacher[tid]
    key = _name_key(teacher["name"])
    ours = await session.scalar(select(Teacher.id).where(or_(
        Teacher.edupage_id == tid, func.lower(Teacher.full_name) == key, func.lower(Teacher.short_name) == key)).limit(1))
    t = translator(get_lang(request))
    return await render(
        request, "admin/timetable_week.html", admin, (cid, campaigns),
        # Opened from the teacher's own page this step reads "Timetable": Teachers › Name › Timetable.
        crumb(plain(teacher["name"]), "tt-teacher", f"teacher:{ours}" if ours else f"tt-teacher:{tid}",
              label(t, "tt_title")),
        [section(t, campaign_url(cid), "/admin/timetable", tab="teachers")], **timetable_context(request, tt),
        kind="teacher", entity=teacher, week=tt.week(lessons, datetime.now(TASHKENT)), stats=tt.stats(lessons),
        people=len({p for x in lessons for p in x["groups"]}), subjects=tt.subjects(lessons, "groups"),
        teachers=sorted(tt.faculty_teachers(), key=lambda x: x["name"].lower()), feedback_id=ours)


@router.post("/timetable/sync")
async def timetable_sync(request: Request, next: str = Form("/admin/timetable"), admin: Admin = Depends(superadmin),
                         session: AsyncSession = Depends(get_session, scope="function")):
    """Download the timetable now (it is also checked automatically every few hours)."""
    t = translator(get_lang(request))
    changed, error = await refresh_from_edupage(session, admin=admin.username)
    await session.commit()
    if error:
        log.warning("Timetable sync failed: %s", error)
        flash(request, "bad", t("tt_sync_error"))
    else:
        flash(request, "good", t("tt_sync_changed") if changed else t("tt_sync_same"))
    return RedirectResponse(next if next.startswith("/admin") else "/admin/timetable", status_code=303)


# ---------------------------------------------------------------- Excel export

def xlsx_response(filename: str, headers: list[str], rows: list) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}_{stamp}.xlsx"'})


@router.get("/export/comments.xlsx")
async def export_comments(request: Request, campaign: str | None = None, q: str | None = None, band: str | None = None,
                          teacher: OptionalId = None, group: OptionalId = None, admin: Admin = Depends(current_admin),
                          session: AsyncSession = Depends(get_session, scope="function")):
    """The comments page as a spreadsheet, with the filters that are set on it."""
    cid, _ = await resolve_campaign(session, campaign)
    t = translator(get_lang(request))
    rows = await an.comment_export_rows(session, cid, q, band, teacher, group, by_student=admin.sees_identity)
    rows = [(*r[:10], float(r[10]) if r[10] is not None else None, r[11]) for r in rows]
    headers = [t("col_survey"), t("col_date"), t("col_teacher"), t("col_department"), t("col_student"), "HEMIS",
               t("col_group"), *criteria_labels(t).values(), t("overall"), t("col_comment")]
    if not admin.sees_identity:  # no student name or HEMIS id in the file
        headers = headers[:4] + headers[6:]
        rows = [r[:4] + r[6:] for r in rows]
    return xlsx_response("comments", headers, rows)
