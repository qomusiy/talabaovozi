from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.faculty import group_key
from app.importers.excel import (ImportFileError, import_assignments, import_group_sizes, import_teachers,
                                 parse_timetable, teaching_sheets, tidy_name)
from app.importers.templates import TEMPLATES, template_xlsx
from app.models import Group, Semester, Teacher, TeachingAssignment
from app.services.teaching import name_key, teacher_name_keys


def xlsx(*rows) -> BytesIO:
    wb = Workbook()
    for r in rows:
        wb.active.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_group_keys_match_hemis_and_edupage_spellings():
    assert group_key("26-01-IQT uz") == group_key("26-01 IQT") == "26-01 IQT"
    assert group_key("26-03 IQT ru") == group_key("26-03 IQT (ru)")
    assert group_key("25-03 IQT O‘zbek") == group_key("25-03 IQT (ing)")
    assert group_key("24-20 FTO‘IT") == group_key("24-20 FTO'IT ") == "24-20 FTOIT"
    assert group_key("26-13a BXA") != group_key("26-13 BXA")
    assert group_key("Iqtisodiyot") is None


def test_names():
    assert tidy_name("TOSHMATOV  ESHMAT O‘G‘LI") == "Toshmatov Eshmat O‘g‘li"
    assert tidy_name("Karimov Aziz") == "Karimov Aziz"
    assert name_key("Abdug‘aniyev A.A") == name_key("abdug'aniyev AA") == "abdug'aniyev aa"
    assert {"qodirov rs", "qodirov rustam", "qodirov r"} <= teacher_name_keys("Qodirov Rustam Sobirovich")


async def test_group_sizes_from_a_hemis_export_count_students_and_keep_no_names(session):
    f = xlsx(
        ["Talabalar ro'yxati"],  # title row above the header, common in real exports
        ["Talaba ID", "Familiya ismi sharifi", "Passport", "Группа", "Talaba kursi"],
        [368211100001.0, "Karimov  Aziz", "as 1234567", "25-16 IQT", "2-kurs"],  # float id
        ["368211100002", "Saidova Madina", "AB7654321", "25-16-IQT uz", "2-kurs"],  # same group, HEMIS spelling
        ["368211100003", "Toshmatov Eshmat", "", "25-17 IQT", "2-kurs"],
        ["368211100002", "Duplicate", "AB7654321", "25-16 IQT", "2-kurs"],         # the same student twice
        ["368211100004", "No Group", "", "", ""],
    )
    report = await import_group_sizes(session, f)
    assert (report.updated, report.skipped, len(report.errors)) == (2, 1, 1)
    sizes = dict((await session.execute(select(Group.key, Group.student_count))).all())
    assert sizes == {"25-16 IQT": 2, "25-17 IQT": 1}
    assert (await session.scalar(select(Group.course).where(Group.key == "25-16 IQT"))) == 2


async def test_group_sizes_from_a_count_column_and_clearing_missing_groups(session):
    await import_group_sizes(session, xlsx(["Guruh", "Talabalar soni"], ["25-01 IQT", 28], ["25-02 IQT", "30"]))
    report = await import_group_sizes(session, xlsx(["Guruh", "Talabalar soni"], ["25-01 IQT", 27], ["25-03 IQT", "x"]))
    assert len(report.errors) == 1 and "not a number" in report.messages()[0]
    report = await import_group_sizes(session, xlsx(["Guruh", "Talabalar soni"], ["25-01 IQT", 26]), clear_missing=True)
    assert report.deactivated == 1
    sizes = dict((await session.execute(select(Group.key, Group.student_count))).all())
    assert sizes == {"25-01 IQT": 26, "25-02 IQT": None}


async def test_missing_column_gives_clear_error(session):
    with pytest.raises(ImportFileError, match="Guruh"):
        await import_group_sizes(session, xlsx(["F.I.Sh.", "Pasport"], ["A", "G"]))


async def test_assignments_match_teacher_by_id_name_or_timetable_name(session):
    await import_teachers(session, xlsx(["O‘qituvchi ID", "F.I.Sh.", "Jadvaldagi ism", "Kafedra", "Xona raqami"],
                                        ["T1", "Qodirov Rustam Sanjarovich", "Qodirov RS", "Iqtisodiyot", "2-14"]))
    sem = Semester(name="S", starts_on=date(2026, 9, 1), ends_on=date(2027, 1, 20), is_current=True)
    session.add(sem)
    report = await import_assignments(session, xlsx(
        ["O‘qituvchi ID", "O‘qituvchi F.I.Sh.", "Guruh", "Fan", "Mashg‘ulot turi"],
        ["T1", "", "25-16 IQT", "Makroiqtisodiyot", "Ma'ruza"],
        ["", "qodirov rustam sanjarovich", "25-16 IQT", "Makroiqtisodiyot", "Amaliy"],
        ["", "Qodirov R.S.", "25-17 IQT", "Makroiqtisodiyot (S)", ""],  # type from the subject suffix
        ["", "", "25-17 IQT", "Makroiqtisodiyot", "Seminar"],
    ), sem)
    assert report.created == 3 and len(report.errors) == 1
    types = set((await session.execute(select(TeachingAssignment.lesson_type))).scalars())
    assert types == {"lecture", "practice", "seminar"}
    assert (await session.scalar(select(func.count(Teacher.id)))) == 1

    # Re-import replaces instead of duplicating.
    await import_assignments(session, xlsx(["O‘qituvchi ID", "Guruh", "Fan"], ["T1", "25-16 IQT", "Fan"]), sem)
    assert (await session.execute(select(func.count()).select_from(TeachingAssignment))).scalar() == 1


def test_timetable_from_excel():
    data, report = parse_timetable(template_xlsx("timetable", sample=True))
    assert report.errors == [] and data["source"] == "excel"
    assert [g["name"] for g in data["groups"]] == ["25-16 IQT", "25-17 IQT"]
    first, second = data["lessons"]
    assert (first["day"], first["period"], first["type"], len(first["groups"])) == (0, 1, "lecture", 2)
    assert (second["week"], second["type"]) == ("A", "practice")


async def test_every_template_imports(session):
    """Templates offered in the dashboard must always be accepted by their importer."""
    session.add(Semester(name="S", starts_on=date(2026, 9, 1), ends_on=date(2027, 1, 20), is_current=True))
    await session.flush()
    sem = await session.scalar(select(Semester))
    assert not (await import_group_sizes(session, template_xlsx("students", sample=True))).errors
    teachers, assignments = teaching_sheets(template_xlsx("teaching", sample=True))
    assert len(teachers) == 2 and len(assignments) == 3
    assert not (await import_teachers(session, None, rows=teachers)).errors
    assert not (await import_assignments(session, None, sem, rows=assignments)).errors
    # The single-purpose importers find their sheet in the two-sheet file too.
    assert not (await import_teachers(session, template_xlsx("teaching", sample=True))).errors
    assert set(TEMPLATES) == {"students", "teaching", "timetable"}


async def test_template_examples_are_never_imported(session):
    """A downloaded template has the example only on its "Namuna" sheet: uploading it unchanged adds nothing."""
    from openpyxl import load_workbook
    wb = load_workbook(template_xlsx("students", {"groups": ["25-16 IQT"]}))
    assert wb.sheetnames[:3] == ["Guruhlar", "Yo‘riqnoma", "Namuna"] and wb.active.title == "Guruhlar"
    assert wb["Guruhlar"]["A2"].value is None
    report = await import_group_sizes(session, template_xlsx("students"))
    assert report.updated == 0
    assert teaching_sheets(template_xlsx("teaching")) == ([], [])
