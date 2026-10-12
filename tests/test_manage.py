"""The dashboard's management pages: surveys, data, bot, admins - and the protections around them."""
import asyncio
import re
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import func, select

from app import demo, settings_store
from app.importers.templates import template_xlsx
from app.models import (Admin, Broadcast, BotUser, Department, Feedback, Group, ImportLog, Semester, SurveyCampaign,
                        Teacher, TeachingAssignment)
from app.security import hash_password
from app.services import broadcast, teaching
from app.web.main import create_app
from app.web.security import throttle
from app.web.telegram_auth import sign_init_data
from tests.test_timetable import edupage_raw
from app.timetable import edupage

RATING = {"teaching": 8, "fairness": 7, "professionalism": 9}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
async def world(session):
    sem = Semester(name="2026–2027 kuzgi", starts_on=date(2026, 9, 7), ends_on=date(2027, 1, 24), is_current=True)
    session.add_all([sem, Admin(username="boss", password_hash=hash_password("secret123"), role="superadmin"),
                     Admin(username="viewer", password_hash=hash_password("secret123"), role="viewer")])
    await session.commit()
    throttle.failures.clear()
    return {"sem": sem}


async def client_for(username: str) -> AsyncClient:
    c = AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test")
    await c.post("/admin/login", data={"username": username, "password": "secret123"})
    return c


@pytest.fixture
async def boss(world):
    c = await client_for("boss")
    yield c
    await c.aclose()


async def with_timetable(session) -> None:
    """The small EduPage timetable from test_timetable, stored -> groups, teachers, assignments."""
    data = edupage.parse(edupage_raw(), "Iqtisodiyot") | {"source": "edupage"}
    await teaching.store_timetable(session, data, title="T", source="https://x/timetable/", admin="test")
    await session.commit()


# ---------------------------------------------------------------- Timetable -> teaching data

async def test_timetable_builds_groups_teachers_and_assignments(session, world):
    await with_timetable(session)
    groups = {g.name: g for g in (await session.execute(select(Group))).scalars()}
    assert set(groups) == {"26-01 IQT", "26-03 IQT (ru)", "25-38 FTO‘IT"}  # only the faculty's groups
    assert groups["26-03 IQT (ru)"].key == "26-03 IQT" and groups["26-03 IQT (ru)"].language == "ru"
    assert groups["25-38 FTO‘IT"].code == "FTO'IT"
    teachers = {t.full_name: t for t in (await session.execute(select(Teacher))).scalars()}
    assert set(teachers) == {"Qodirov RS", "Saidova Madina"} and teachers["Qodirov RS"].edupage_id == "n1"
    rows = (await session.execute(select(Teacher.full_name, Group.name, TeachingAssignment.lesson_type)
                                  .join(Teacher).join(Group))).all()
    # The lecture is for two groups; English ("Ingliz tili") is taught to another faculty's group.
    assert sorted(rows) == [("Qodirov RS", "26-01 IQT", "lecture"), ("Qodirov RS", "26-03 IQT (ru)", "lecture"),
                            ("Saidova Madina", "26-01 IQT", "practice")]
    assert (await session.scalar(select(ImportLog.kind))) == "timetable"

    # The same timetable again changes nothing; students imported with HEMIS spelling land in the same group.
    data = edupage.parse(edupage_raw(), "Iqtisodiyot") | {"source": "edupage"}
    assert not await teaching.store_timetable(session, data, title="T", source="x", admin="test")

    # Excluded subjects are not rated.
    await settings_store.save(session, excluded_subjects=["makroiqtisodiyot"])
    assert await teaching.rebuild_assignments(session) == 0


# ---------------------------------------------------------------- Kafedras

def test_kafedra_is_guessed_from_subjects():
    from app.services.departments import guess
    assert guess([("Makroiqtisodiyot", ["26-01 IQT"])]) == "Iqtisodiyot"
    assert guess([("Ehtimollar nazariyasi va matematik statistika", [])]) == "Oliy matematika"
    assert guess([("Statistika", []), ("STАTА dasturi", [])]) == "Buxgalteriya hisobi va audit"  # Cyrillic "А"s
    assert guess([("Iqtisodiyotda axborot kommunikatsion texnologiyalar", [])]) == "Axborot texnologiyalari"
    # Literature is taught in both philology specialties: the group decides.
    assert guess([("Adabiyotshunoslik", ["25-38 FTO‘RT"])]) == "O‘zbek tili va adabiyoti"
    assert guess([("Adabiyotshunoslik", ["25-38 FTO‘IT"])]) == "Ingliz tili"
    assert guess([("Sportda xorijiy til", []), ("O‘zbekistonning eng yangi tarixi", []),
                  ("O‘zbekistonning eng yangi tarixi", [])]) == "Gumanitar fanlar"
    assert guess([("Kelajak soati", [])]) is None


async def test_timetable_teachers_get_a_kafedra(session, world):
    await with_timetable(session)
    rows = dict((await session.execute(select(Teacher.full_name, Department.name).join(Department))).all())
    assert rows == {"Qodirov RS": "Iqtisodiyot", "Saidova Madina": "Iqtisodiyot"}  # Makroiqtisodiyot
    assert (await session.scalar(select(Department.faculty))) == "Iqtisodiyot"


# ---------------------------------------------------------------- Editing data by hand

async def test_edit_teachers_and_kafedras(boss, session):
    await with_timetable(session)
    for tab in ("teachers", "departments", "groups"):
        assert (await boss.get(f"/admin/records?tab={tab}")).status_code == 200, tab
    teacher = await session.scalar(select(Teacher).where(Teacher.full_name == "Qodirov RS"))
    r = await boss.post("/admin/records/departments/new", data={"name": "Oliy matematika", "faculty": "Yer resurslari va kadastr"})
    assert r.status_code == 303
    math = await session.scalar(select(Department.id).where(Department.name == "Oliy matematika"))
    await boss.post(f"/admin/records/teachers/{teacher.id}", data={
        "full_name": "Qodirov  Rustam Sobirovich", "department": str(math), "position": "dotsent", "next": "/admin/records"})
    await session.refresh(teacher)
    assert (teacher.full_name, teacher.department_id, teacher.position, teacher.is_active) == (
        "Qodirov Rustam Sobirovich", math, "dotsent", False)                    # "active" not ticked
    assert "Yer resurslari va kadastr fakulteti" in (await boss.get("/admin/teachers")).text  # filter grouped by faculty
    # The next timetable sync keeps what was set by hand.
    await teaching.rebuild_assignments(session)
    await session.commit()
    await session.refresh(teacher)
    assert (teacher.full_name, teacher.department_id, teacher.is_active) == ("Qodirov Rustam Sobirovich", math, False)
    # A kafedra name can't be taken twice; deleting one leaves its teachers without a kafedra.
    await boss.post("/admin/records/departments/new", data={"name": "oliy matematika"})
    assert await session.scalar(select(func.count(Department.id)).where(func.lower(Department.name) == "oliy matematika")) == 1
    await boss.post(f"/admin/records/departments/{math}/delete", data={})
    await session.refresh(teacher)
    assert teacher.department_id is None
    # ...and "fill in from subjects" gives it back the guessed one.
    await boss.post("/admin/records/departments-guess", data={})
    await session.refresh(teacher)
    assert (await session.get(Department, teacher.department_id)).name == "Iqtisodiyot"


async def test_edit_groups_and_their_sizes(boss, session, world):
    await with_timetable(session)
    timetable_group = await session.scalar(select(Group).where(Group.name == "26-01 IQT"))
    before = (timetable_group.name, timetable_group.course, timetable_group.language)
    await boss.post(f"/admin/records/groups/{timetable_group.id}",
                    data={"name": "Renamed", "course": "4", "language": "ru", "students": "27"})
    await session.refresh(timetable_group)
    assert (timetable_group.name, timetable_group.course, timetable_group.language) == before  # the timetable's
    assert timetable_group.student_count == 27                                                # but the size is ours
    await boss.post("/admin/records/groups/new", data={"name": "26-99 BXA", "course": "1", "students": "20"})
    group = await session.scalar(select(Group).where(Group.name == "26-99 BXA"))
    assert (group.key, group.code, group.course, group.student_count) == ("26-99 BXA", "BXA", 1, 20)
    await boss.post(f"/admin/records/groups/{group.id}", data={"name": "26-98 BXA (ru)", "course": "2", "language": "ru"})
    await session.refresh(group)
    assert (group.name, group.key, group.course, group.language, group.student_count) == (
        "26-98 BXA (ru)", "26-98 BXA", 2, "ru", None)
    r = await boss.post(f"/admin/records/groups/{group.id}", data={"name": "26-98 BXA (ru)", "students": "many"})
    assert "0 dan 1000 gacha" in (await boss.get(r.headers["location"])).text
    assert "26-98 BXA (ru)" in (await boss.get("/admin/records?tab=groups")).text

    # A group that has ratings can't be deleted; a teacher who has them neither.
    exists = lambda model, id_: session.scalar(select(model.id).where(model.id == id_))  # noqa: E731
    teacher = await session.scalar(select(Teacher).where(Teacher.full_name == "Qodirov RS"))
    camp = SurveyCampaign(semester=world["sem"], title="R", opens_at=datetime.now(timezone.utc) - timedelta(days=1),
                          closes_at=datetime.now(timezone.utc) + timedelta(days=1))
    session.add_all([camp, BotUser(id=41)])
    await session.flush()
    session.add(Feedback(campaign=camp, user_id=41, teacher=teacher, group_id=group.id, teaching=8, fairness=8,
                         professionalism=8))
    await session.commit()
    await boss.post(f"/admin/records/teachers/{teacher.id}/delete", data={})
    await boss.post(f"/admin/records/groups/{group.id}/delete", data={})
    assert await exists(Teacher, teacher.id) and await exists(Group, group.id)
    await session.execute(__import__("sqlalchemy").delete(Feedback))
    await session.commit()
    await boss.post(f"/admin/records/groups/{group.id}/delete", data={})
    assert not await exists(Group, group.id)                                # no ratings now: deleted


async def test_group_numbers_students_know(boss, session, world):
    await with_timetable(session)
    g1 = await session.scalar(select(Group).where(Group.name == "26-01 IQT"))
    await boss.post(f"/admin/records/groups/{g1.id}", data={"number": " 302 ", "students": "25"})  # a timetable group too
    await session.refresh(g1)
    assert (g1.number, g1.student_count) == ("302", 25)
    other = Group(name="26-77 BXA", course=1)
    session.add(other)
    await session.commit()
    r = await boss.post(f"/admin/records/groups/{other.id}", data={"name": "26-77 BXA", "number": "302"})
    assert "302" in (await boss.get(r.headers["location"])).text and (await session.get(Group, other.id)).number is None

    # The analysis pages show it only for whoever switches it on (Ko‘rinish).
    assert "302 · 26-01 IQT" not in (await boss.get("/admin/groups")).text
    await boss.post("/admin/account/display", data={"group_numbers": "true"})
    assert "302 · 26-01 IQT" in (await boss.get("/admin/groups")).text

    # The group sizes file can carry numbers too; a row with only a number leaves the size alone.
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Guruh", "Raqami"])
    wb.active.append(["26-01 IQT", "301"])
    wb.active.append(["26-77 BXA", "301"])            # taken by the row above: an error, not a duplicate
    buf = BytesIO()
    wb.save(buf)
    from app.importers import excel
    report = await excel.import_group_sizes(session, BytesIO(buf.getvalue()))
    await session.refresh(g1)
    assert (g1.number, g1.student_count) == ("301", 25) and [e[1] for e in report.errors] == ["number_taken"]


# ---------------------------------------------------------------- Access & security

async def test_viewers_only_see_analytics(world):
    c = await client_for("viewer")
    r = await c.get("/admin")
    assert r.status_code == 200 and "/admin/surveys" not in r.text and "/admin/data" not in r.text
    for url in ["/admin/surveys", "/admin/data", "/admin/bot", "/admin/admins", "/admin/records"]:
        assert (await c.get(url)).status_code == 403, url
    assert (await c.post("/admin/records/departments/new", data={"name": "X"})).status_code == 403
    assert (await c.post("/admin/bot/settings", data={})).status_code == 403
    assert (await c.get("/admin/account")).status_code == 200
    await c.aclose()


async def test_login_is_blocked_after_five_wrong_passwords(world):
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        for _ in range(5):
            assert (await c.post("/admin/login", data={"username": "boss", "password": "nope"})).status_code == 401
        r = await c.post("/admin/login", data={"username": "boss", "password": "secret123"})
        assert r.status_code == 429 and "daqiqadan keyin" in r.text


async def test_forms_cannot_be_posted_from_other_sites(boss):
    r = await boss.post("/admin/bot/settings", data={}, headers={"Origin": "https://evil.example"})
    assert r.status_code == 403
    r = await boss.get("/admin")
    assert r.headers["x-frame-options"] == "DENY" and r.headers["cache-control"] == "no-store"


async def test_every_management_page_renders(boss, session):
    await with_timetable(session)
    for url in ["/admin/surveys", "/admin/data", "/admin/bot", "/admin/admins", "/admin/account",
                "/admin/timetable", "/admin/records?tab=groups"]:
        assert (await boss.get(url)).status_code == 200, url
    for lang in ("en", "ru"):
        boss.cookies.set("admin_lang", lang)
        for url in ["/admin/surveys", "/admin/data", "/admin/bot", "/admin/admins"]:
            assert (await boss.get(url)).status_code == 200, (lang, url)


# ---------------------------------------------------------------- Survey rounds

async def test_survey_rounds(boss, session):
    now = datetime.now(timezone.utc) + timedelta(hours=5)  # Tashkent wall clock, as the form sends it
    fmt = lambda d: d.strftime("%Y-%m-%dT%H:%M")  # noqa: E731
    r = await boss.post("/admin/surveys/new", data={"title": "Oraliq", "opens": fmt(now - timedelta(hours=1)),
                                                     "closes": fmt(now + timedelta(days=7))}, follow_redirects=False)
    assert r.status_code == 303
    c = await session.scalar(select(SurveyCampaign))
    assert c.title == "Oraliq" and c.opens_at < datetime.now(timezone.utc) < c.closes_at

    # Overlapping or backwards rounds are refused.
    await boss.post("/admin/surveys/new", data={"title": "Clash", "opens": fmt(now + timedelta(days=1)),
                                                 "closes": fmt(now + timedelta(days=9))})
    assert "ustma-ust" in (await boss.get("/admin/surveys")).text
    await boss.post("/admin/surveys/new", data={"title": "Back", "opens": fmt(now + timedelta(days=30)),
                                                 "closes": fmt(now + timedelta(days=29))})
    assert (await session.scalar(select(func.count(SurveyCampaign.id)))) == 1

    await boss.post(f"/admin/surveys/{c.id}/close")
    await session.refresh(c)
    assert c.closes_at <= datetime.now(timezone.utc)
    await boss.post(f"/admin/surveys/{c.id}/edit", data={"title": "Renamed", "opens": fmt(now + timedelta(days=1)),
                                                          "closes": fmt(now + timedelta(days=5))})
    await session.refresh(c)
    assert c.title == "Renamed" and c.opens_at > datetime.now(timezone.utc)
    await boss.post(f"/admin/surveys/{c.id}/open")
    await session.refresh(c)
    assert c.opens_at <= datetime.now(timezone.utc)
    await boss.post(f"/admin/surveys/{c.id}/delete")
    assert (await session.scalar(select(func.count(SurveyCampaign.id)))) == 0


async def test_round_times_compare_to_the_minute_and_stay_in_the_semester(boss, session):
    """Closing one round and opening the next "now" leaves seconds the form doesn't show; saving the next round
    unchanged must not count as overlapping. A round can't run past its semester."""
    now = datetime.now(timezone.utc) + timedelta(hours=5)
    fmt = lambda d: d.strftime("%Y-%m-%dT%H:%M")  # noqa: E731
    await boss.post("/admin/surveys/new", data={"title": "F1", "opens": fmt(now - timedelta(hours=2)),
                                                 "closes": fmt(now + timedelta(days=3))})
    f1 = await session.scalar(select(SurveyCampaign).where(SurveyCampaign.title == "F1"))
    await boss.post(f"/admin/surveys/{f1.id}/close")                      # closes_at = now, with seconds
    await boss.post("/admin/surveys/new", data={"title": "f2", "opens": fmt(now + timedelta(days=1)),
                                                 "closes": fmt(now + timedelta(days=4))})
    f2 = await session.scalar(select(SurveyCampaign).where(SurveyCampaign.title == "f2"))
    await boss.post(f"/admin/surveys/{f2.id}/open")                       # opens_at = now, with seconds
    await session.refresh(f2)
    form = await boss.get("/admin/surveys")
    opens = re.search(rf'action="/admin/surveys/{f2.id}/edit".*?name="opens" value="([^"]+)"', form.text, re.S).group(1)
    await boss.post(f"/admin/surveys/{f2.id}/edit", data={"title": "f2 renamed", "opens": opens,
                                                           "closes": fmt(now + timedelta(days=10))})
    await session.refresh(f2)
    assert f2.title == "f2 renamed"                                         # saved: no false overlap with F1

    r = await boss.post(f"/admin/surveys/{f2.id}/edit", data={"title": "f2", "opens": opens, "closes": "2028-05-15T20:00"},
                        follow_redirects=False)
    assert r.headers["location"] == f"/admin/surveys?edit={f2.id}"       # the form stays open to fix
    page = (await boss.get(r.headers["location"])).text
    assert "semestr" in page.lower() and "24.01.2027" in page
    await session.refresh(f2)
    assert f2.closes_at.year == now.year                                    # not saved


async def test_demo_data_is_created_over_real_data_and_removed_cleanly(boss, session):
    await with_timetable(session)
    group = await session.scalar(select(Group).where(Group.name == "26-01 IQT"))
    group.student_count = 40
    session.add(BotUser(id=5, group_id=group.id, onboarded_at=datetime.now(timezone.utc)))   # a real account
    await session.commit()
    result = await demo.create_demo(session)
    await session.commit()
    assert result["rounds"] == 3 and result["feedback"] > 0 and 0 < result["students"] <= 40 + demo.DEFAULT_GROUP_SIZE
    assert "Demo" in (await boss.get("/admin")).text  # the banner
    assert not await session.scalar(select(func.count(Feedback.id)).where(Feedback.user_id == 5))
    assert await session.scalar(select(func.max(BotUser.id)).where(BotUser.is_demo)) < 0   # never a real account

    await boss.post("/admin/surveys/demo/clear")
    assert (await session.scalar(select(func.count(Feedback.id)))) == 0
    assert (await session.scalar(select(func.count(BotUser.id)))) == 1                   # only the real one


# ---------------------------------------------------------------- Data

async def test_uploads_and_templates(boss, session):
    r = await boss.get("/admin/data/template/students.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert (await boss.get("/admin/data/template/nope.xlsx")).status_code == 404

    files = {"file": ("students.xlsx", template_xlsx("students", sample=True).read(), XLSX)}
    r = await boss.post("/admin/data/students", files=files, follow_redirects=False)
    assert r.status_code == 303
    assert "2 ta yangilandi" in (await boss.get("/admin/data")).text
    assert dict((await session.execute(select(Group.name, Group.student_count))).all()) == {
        "25-16 IQT": 28, "26-03 IQT (ru)": 25}
    assert "53 talaba, 2 ta guruhda" in (await boss.get("/admin/data")).text
    assert (await boss.get("/admin/data/temporary-codes.xlsx")).status_code == 404

    bad = {"file": ("x.xlsx", b"not a spreadsheet", XLSX)}
    await boss.post("/admin/data/students", files=bad)
    assert ".xlsx" in (await boss.get("/admin/data")).text

    # One file: teacher details + who teaches whom. Assignments from Excel switch the source (the timetable no
    # longer overwrites them).
    files = {"file": ("t.xlsx", template_xlsx("teaching", sample=True).read(), XLSX)}
    await boss.post("/admin/data/teaching", files=files)
    assert await settings_store.get(session, "assignments_source") == "excel"
    assert (await session.scalar(select(func.count(TeachingAssignment.id)))) == 3
    assert await session.scalar(select(Teacher.position).where(Teacher.full_name == "Qodirov Rustam Sobirovich")) == "dotsent"
    files = {"file": ("t.xlsx", template_xlsx("teaching").read(), XLSX)}           # unfilled template
    await boss.post("/admin/data/teaching", files=files)
    assert "to‘ldirilgan qator yo‘q" in (await boss.get("/admin/data")).text

    files = {"file": ("tt.xlsx", template_xlsx("timetable", sample=True).read(), XLSX)}
    await boss.post("/admin/data/timetable", files=files)
    values = await settings_store.load(session, fresh=True)
    assert values["timetable_auto_sync"] is False  # manual timetable: stop EduPage checks
    assert "tt.xlsx" in (await boss.get("/admin/data")).text


async def test_new_semester_and_sync_settings(boss, session, world):
    await with_timetable(session)
    await boss.post("/admin/data/semester", data={"name": "2026–2027 bahorgi", "starts": "2027-02-01",
                                                  "ends": "2027-06-15"})
    current = await session.scalar(select(Semester).where(Semester.is_current))
    assert current.name == "2026–2027 bahorgi"
    assert (await session.scalar(select(func.count(TeachingAssignment.id)).where(
        TeachingAssignment.semester_id == current.id))) == 3  # rebuilt for the new semester
    await boss.post("/admin/data/settings", data={"timetable_sync_hours": "12", "assignments_source": "timetable",
                                                  "excluded_subjects": "Kelajak soati\nMakroiqtisodiyot"})
    values = await settings_store.load(session, fresh=True)
    assert values["timetable_auto_sync"] is False and values["timetable_sync_hours"] == 12
    assert (await session.scalar(select(func.count(TeachingAssignment.id)).where(
        TeachingAssignment.semester_id == current.id))) == 0


# ---------------------------------------------------------------- Bot & survey settings

async def test_bot_settings_change_the_survey(boss, session):
    await with_timetable(session)
    group = await session.scalar(select(Group).where(Group.name == "26-01 IQT"))
    teacher = await session.scalar(select(Teacher).where(Teacher.full_name == "Qodirov RS"))
    st = BotUser(id=555, group_id=group.id, onboarded_at=datetime.now(timezone.utc))
    now = datetime.now(timezone.utc)
    sem = await session.scalar(select(Semester).where(Semester.is_current))
    session.add_all([st, SurveyCampaign(semester=sem, title="R", opens_at=now - timedelta(days=1),
                                        closes_at=now + timedelta(days=1))])
    await session.commit()
    await boss.post("/admin/bot/settings", data={"survey_notice": "Deadline: Friday", "support_contact": "@help"})
    tg = {"X-Telegram-Init-Data": sign_init_data({"id": 555}, "123:TEST")}
    data = (await boss.get("/api/survey", headers=tg)).json()
    assert data["settings"] == {"comments": False, "edit": False, "notice": "Deadline: Friday", "comment_required": False}

    url = f"/api/survey/teachers/{teacher.id}"
    assert (await boss.post(url, json={**RATING, "comment": "hi"}, headers=tg)).status_code == 200
    assert (await session.scalar(select(Feedback.comment))) is None  # comments are off
    r = await boss.post(url, json=RATING, headers=tg)
    assert r.status_code == 400 and r.json()["detail"] == "locked"  # editing is off
    assert await settings_store.get(session, "support_contact") == "@help"

    # A required comment: a rating without one (at least 10 characters) is refused.
    await boss.post("/admin/bot/settings", data={"survey_comments": "true", "survey_comment_required": "true",
                                                 "survey_edit": "true"})
    settings_store.reset_cache()
    assert (await boss.get("/api/survey", headers=tg)).json()["settings"]["comment_required"] is True
    r = await boss.post(url, json={**RATING, "comment": " ok "}, headers=tg)
    assert r.status_code == 400 and r.json()["detail"] == "comment_required"
    assert (await boss.post(url, json={**RATING, "comment": "Mavzuni tushunarli tushuntiradi"}, headers=tg)).status_code == 200


class FakeBot:
    """Stands in for the bot when the dashboard sends messages; remembers what went where."""
    sent: list = []

    def __init__(self, *a, **kw):
        self.session = self

    async def send_message(self, chat_id, text, reply_markup=None):
        FakeBot.sent.append((chat_id, "text", text))

    async def send_photo(self, chat_id, photo, caption=None, reply_markup=None):
        FakeBot.sent.append((chat_id, "photo", caption))
        from types import SimpleNamespace
        return SimpleNamespace(photo=[SimpleNamespace(file_id="PHOTO-1")])

    async def close(self):
        pass


def png_bytes() -> bytes:
    from io import BytesIO
    from PIL import Image
    out = BytesIO()
    Image.new("RGB", (40, 20), "#5b5bd6").save(out, "PNG")
    return out.getvalue()


async def test_broadcast_reaches_the_chosen_users(boss, session, monkeypatch, tmp_path):
    await with_timetable(session)
    group = await session.scalar(select(Group).where(Group.name == "26-01 IQT"))
    other = await session.scalar(select(Group).where(Group.name == "26-03 IQT (ru)"))
    now = datetime.now(timezone.utc)
    session.add_all([BotUser(id=11, group_id=group.id, onboarded_at=now),
                     BotUser(id=12, group_id=other.id, onboarded_at=now),
                     BotUser(id=13, group_id=group.id),                               # hasn't finished the intro
                     BotUser(id=14, group_id=group.id, onboarded_at=now, blocked=True),
                     BotUser(id=-15, group_id=group.id, onboarded_at=now, is_demo=True)])  # demo: no real account
    # Saidova teaches only 26-01 IQT: her kafedra's users are that group's.
    dept = Department(name="Buxgalteriya")
    session.add(dept)
    saidova = await session.scalar(select(Teacher).where(Teacher.full_name == "Saidova Madina"))
    saidova.department = dept
    await session.commit()

    async def n(**params):
        return (await boss.get("/admin/bot/audience", params=params)).json()["n"]
    assert await n(audience="all") == 2
    assert await n(audience="group", values=[group.id]) == 1
    assert await n(audience="group", values=[group.id, other.id]) == 2
    assert await n(audience="group") == 0                       # nothing chosen yet
    assert await n(audience="department", values=[dept.id]) == 1
    assert await n(audience="specialty", values=["IQT"]) == 2
    assert await n(audience="all", survey="unfinished") == 0   # no open survey: nobody to remind

    FakeBot.sent = []
    monkeypatch.setattr(broadcast.chatlog, "new_bot", FakeBot)
    monkeypatch.setattr(broadcast, "BROADCAST_MEDIA", tmp_path)
    await boss.post("/admin/bot/broadcast", data={"text": "Salom!", "audience": "group", "values": [str(group.id)]})
    await asyncio.gather(*broadcast._running)
    assert FakeBot.sent == [(11, "text", "Salom!")]
    b = await session.scalar(select(Broadcast))
    await session.refresh(b)
    assert (b.status, b.sent, b.total) == ("done", 1, 1)
    page = (await boss.get("/admin/bot")).text
    assert "Salom!" in page and "26-01 IQT" in page            # the history says who got it

    # With a photo: uploaded once, the text under it; to a kafedra.
    FakeBot.sent = []
    await boss.post("/admin/bot/broadcast", data={"text": "Rasm bilan", "audience": "department", "values": [str(dept.id)]},
                    files={"media": ("poster.png", png_bytes(), "image/png")})
    await asyncio.gather(*broadcast._running)
    assert FakeBot.sent == [(11, "photo", "Rasm bilan")]
    b = await session.scalar(select(Broadcast).order_by(Broadcast.id.desc()))
    assert b.media_kind == "photo" and (tmp_path / b.media_name).is_file()

    # "Send it to me first" needs the admin's own Telegram, linked on "My account".
    r = (await boss.post("/admin/bot/broadcast/test", data={"text": "Sinov"})).json()
    assert not r["ok"]
    me = await session.scalar(select(Admin).where(Admin.username == "boss"))
    me.telegram_id = 99
    await session.commit()
    FakeBot.sent = []
    r = (await boss.post("/admin/bot/broadcast/test", data={"text": "Sinov"})).json()
    assert r["ok"] and FakeBot.sent == [(99, "text", "Sinov")]
    assert await session.scalar(select(func.count()).select_from(Broadcast)) == 2   # a test isn't a sent message


# ---------------------------------------------------------------- Admins & account

async def test_admins_and_password(boss, session):
    await boss.post("/admin/admins/new", data={"username": "dean", "password": "longpassword", "role": "viewer"})
    dean = await session.scalar(select(Admin).where(Admin.username == "dean"))
    assert dean.role == "viewer"
    await boss.post(f"/admin/admins/{dean.id}/edit", data={"role": "superadmin", "is_active": "true"})
    await session.refresh(dean)
    assert dean.role == "superadmin"
    me = await session.scalar(select(Admin).where(Admin.username == "boss"))
    await boss.post(f"/admin/admins/{me.id}/delete")
    assert await session.get(Admin, me.id) is not None  # can't delete yourself
    await boss.post(f"/admin/admins/{dean.id}/delete")
    assert await session.scalar(select(Admin).where(Admin.username == "dean")) is None

    await boss.post("/admin/account/password", data={"current": "wrong", "new": "newsecret1", "repeat": "newsecret1"})
    await boss.post("/admin/account/password", data={"current": "secret123", "new": "newsecret1", "repeat": "newsecret1"})
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        r = await c.post("/admin/login", data={"username": "boss", "password": "newsecret1"}, follow_redirects=False)
        assert r.status_code == 303


# ---------------------------------------------------------------- Programmer, activity log, viewers without names

@pytest.fixture
async def dev(world, session):
    session.add(Admin(username="agentsmith", password_hash=hash_password("secret123"), role="programmer"))
    await session.commit()
    c = await client_for("agentsmith")
    yield c
    await c.aclose()


async def test_programmer_account_is_hidden_from_other_admins(boss, dev, session):
    from app.models import Admin as A
    smith = await session.scalar(select(A).where(A.username == "agentsmith"))
    page = (await boss.get("/admin/admins")).text
    assert "agentsmith" not in page and "boss" in page
    assert (await boss.post(f"/admin/admins/{smith.id}/delete")).status_code == 404
    assert (await boss.post(f"/admin/admins/{smith.id}/edit", data={"is_active": "true"})).status_code == 404
    for url in ["/admin/audit", "/admin/records?tab=feedback"]:
        r = await boss.get(url)
        assert r.status_code == 404 or "rd_tab_feedback" not in r.text
    assert (await boss.get("/admin/audit")).status_code == 404
    # The programmer sees everything a superadmin does, plus the log, ratings and the data structure.
    assert "agentsmith" in (await dev.get("/admin/admins")).text
    assert (await dev.get("/admin/audit")).status_code == 200
    data = (await dev.get("/admin/data")).text
    assert 'id="schema-data"' in data and "teaching_assignments" in data
    assert 'id="schema-data"' not in (await boss.get("/admin/data")).text


async def test_activity_log_records_logins_pages_and_changes(boss, dev, session):
    from app.models import AuditLog
    await with_timetable(session)
    teacher = await session.scalar(select(Teacher).where(Teacher.full_name == "Qodirov RS"))
    await boss.get("/admin/teachers")
    await boss.post(f"/admin/records/teachers/{teacher.id}", data={"full_name": "Qodirov Rustam", "position": "dotsent",
                                                                   "active": "true", "department": ""})
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        await c.post("/admin/login", data={"username": "boss", "password": "wrong-one"})
    rows = list((await session.execute(select(AuditLog).order_by(AuditLog.id))).scalars())
    kinds = [(r.username, r.kind) for r in rows]
    assert ("boss", "login") in kinds and ("boss", "view") in kinds and ("boss", "action") in kinds
    assert ("boss", "login_failed") in kinds
    failed = next(r for r in rows if r.kind == "login_failed")
    assert "wrong-one" not in str(failed.details)                          # passwords are never stored
    view = next(r for r in rows if r.kind == "view" and r.path == "/admin/teachers")
    assert view.title == "O‘qituvchilar" and view.sid
    action = next(r for r in rows if r.kind == "action" and "/records/teachers/" in r.path)
    assert action.outcome.startswith("Saqlandi") and action.sid == view.sid
    changes = {(c["field"], c["old"], c["new"]) for c in action.details["changes"] if "field" in c}
    assert ("full_name", "Qodirov RS", "Qodirov Rustam") in changes and ("position", None, "dotsent") in changes
    assert any(f == "department_id" and str(old).startswith("Iqtisodiyot #") and new is None for f, old, new in changes)
    page = (await dev.get("/admin/audit?tab=events")).text
    assert "Qodirov Rustam" in page and "Saqlandi" in page
    assert "boss" in (await dev.get("/admin/audit")).text                    # the session list


async def rated_by_aziz(session, world, comment="Zo‘r dars"):
    """A rating by the account of "Aziz Karimov" (@aziz_k, Telegram id 5551234)."""
    await with_timetable(session)
    teacher = await session.scalar(select(Teacher).where(Teacher.full_name == "Qodirov RS"))
    group = await session.scalar(select(Group).where(Group.name == "26-01 IQT"))
    camp = SurveyCampaign(semester=world["sem"], title="R", opens_at=datetime.now(timezone.utc) - timedelta(days=1),
                          closes_at=datetime.now(timezone.utc) + timedelta(days=1))
    user = BotUser(id=5551234, first_name="Aziz", last_name="Karimov", username="aziz_k", group_id=group.id,
                   onboarded_at=datetime.now(timezone.utc), profile={"chat": {"bio": "talaba"}})
    session.add_all([camp, user])
    await session.flush()
    fb = Feedback(campaign=camp, user_id=user.id, teacher=teacher, group_id=group.id, teaching=3, fairness=4,
                  professionalism=5, comment=comment)
    session.add(fb)
    await session.commit()
    return user, teacher, camp, fb


async def test_programmer_can_change_and_delete_a_rating(boss, dev, session, world):
    user, teacher, camp, fb = await rated_by_aziz(session, world)
    assert "Aziz Karimov" in (await dev.get("/admin/records?tab=feedback&q=aziz_k")).text
    assert "Aziz Karimov" in (await dev.get("/admin/records?tab=feedback&q=5551234")).text
    assert (await boss.post(f"/admin/records/feedback/{fb.id}", data={"teaching": "9", "fairness": "9",
                                                                       "professionalism": "9"})).status_code == 404
    await dev.post(f"/admin/records/feedback/{fb.id}", data={"teaching": "9", "fairness": "8", "professionalism": "7",
                                                              "comment": "  fixed  "})
    await session.refresh(fb)
    assert (fb.teaching, fb.fairness, fb.professionalism, fb.comment) == (9, 8, 7, "fixed")
    await dev.post(f"/admin/records/feedback/{fb.id}", data={"teaching": "11", "fairness": "8", "professionalism": "7"})
    await session.refresh(fb)
    assert fb.teaching == 9                                                  # out of range: refused
    await dev.post(f"/admin/records/feedback/{fb.id}/delete")
    assert not await session.scalar(select(Feedback.id).where(Feedback.id == fb.id))


async def test_only_programmers_see_who_rated(boss, dev, session, world):
    """Anonymous: viewers and superadmins see ratings and comments with the group only - also in Excel; programmer
    accounts see the Telegram account (to find abuse)."""
    from io import BytesIO
    from openpyxl import load_workbook
    user, teacher, camp, fb = await rated_by_aziz(session, world)
    for c in (await client_for("viewer"), boss):
        for url in ["/admin", "/admin/comments", f"/admin/teachers/{teacher.id}"]:
            page = (await c.get(url)).text
            assert "Aziz" not in page and "aziz_k" not in page and "Zo‘r dars" in page and "26-01 IQT" in page, url
        assert "Zo‘r dars" not in (await c.get("/admin/comments?q=aziz_k")).text   # can't search by who
        sheet = load_workbook(BytesIO((await c.get("/admin/export/comments.xlsx")).content)).active
        header = [cell.value for cell in sheet[1]]
        assert "Telegram ID" not in header and "Kim baholadi" not in header and sheet.max_row == 2
        assert (await c.get("/admin/users")).status_code == 404
    for url in ["/admin", "/admin/comments", f"/admin/teachers/{teacher.id}"]:
        assert "/admin/users/5551234" in (await dev.get(url)).text, url
    assert "Zo‘r dars" in (await dev.get("/admin/comments?q=aziz_k")).text
    sheet = load_workbook(BytesIO((await dev.get("/admin/export/comments.xlsx")).content)).active
    assert [cell.value for cell in sheet[1]][4:6] == ["Kim baholadi", "Telegram ID"]
    assert [cell.value for cell in sheet[2]][4:6] == ["Aziz Karimov @aziz_k", 5551234]


async def test_users_page_block_suspend_warn_and_delete(boss, dev, session, world, monkeypatch):
    from app.web.admin import users as users_page
    user, teacher, camp, fb = await rated_by_aziz(session, world)
    session.add(BotUser(id=-3, is_demo=True))
    await session.commit()
    html = (await dev.get("/admin/users")).text
    assert "Aziz Karimov" in html and "@aziz_k" in html and "Demo #3" not in html   # demo accounts: their own tab
    assert "Demo #3" in (await dev.get("/admin/users?f=demo")).text
    assert "Aziz" in (await dev.get("/admin/users?q=5551234")).text
    assert "Aziz" not in (await dev.get("/admin/users?f=stopped")).text

    async def no_telegram(u):
        return None
    monkeypatch.setattr(users_page, "refresh_from_telegram", no_telegram)
    page = (await dev.get(f"/admin/users/{user.id}")).text
    assert "Zo‘r dars" in page and "Qodirov RS" in page and "26-01 IQT" in page and "talaba" in page  # bio from getChat

    await dev.post(f"/admin/users/{user.id}/block", data={"blocked": "true", "note": "spam"})
    await session.refresh(user)
    assert user.blocked and user.note == "spam"
    assert "Aziz" in (await dev.get("/admin/users?f=stopped")).text
    await dev.post(f"/admin/users/{user.id}/block", data={})
    await dev.post(f"/admin/users/{user.id}/suspend", data={"hours": "24"})
    await session.refresh(user)
    assert not user.blocked and user.suspended_until > datetime.now(timezone.utc) + timedelta(hours=23)
    await dev.post(f"/admin/users/{user.id}/suspend", data={"hours": "0"})
    await session.refresh(user)
    assert user.suspended_until is None

    FakeBot.sent = []
    monkeypatch.setattr(users_page.chatlog, "new_bot", FakeBot)
    await dev.post(f"/admin/users/{user.id}/warn", data={"text": "Spam qilmang!"})
    await session.refresh(user)
    assert FakeBot.sent == [(5551234, "text", "Spam qilmang!")] and user.warnings == 1

    await dev.post(f"/admin/users/{user.id}/group", data={"group": ""})
    await session.refresh(user)
    assert user.group_id is None
    assert (await boss.post(f"/admin/users/{user.id}/delete")).status_code == 404    # programmer accounts only
    await dev.post(f"/admin/users/{user.id}/ratings/delete", data={"campaign": str(camp.id)})
    assert not await session.scalar(select(Feedback.id).where(Feedback.user_id == user.id))
    await dev.post(f"/admin/users/{user.id}/delete")
    assert await session.get(BotUser, user.id, populate_existing=True) is None


async def test_dashboard_online_switch_and_public_address(boss, session):
    tunnel = {"Cf-Connecting-Ip": "1.2.3.4"}
    assert (await boss.get("/admin/login", headers=tunnel)).status_code == 404       # off by default
    assert (await boss.get("/healthz", headers=tunnel)).text == "ok"
    await boss.post("/admin/bot/online", data={"admin_online": "true"})
    settings_store.reset_cache()
    assert (await boss.get("/admin/login", headers=tunnel)).status_code == 200
    page = (await boss.get("/admin/bot")).text
    assert "Yangi manzil olish" in page and "Admin saytini internetda ochish" in page
    await boss.post("/admin/bot/online", data={})
    settings_store.reset_cache()
    assert (await boss.get("/admin/login", headers=tunnel)).status_code == 404


async def test_reminders_go_once_at_3_days_and_1_day(session, world):
    from app.bot import reminders
    now = datetime.now(timezone.utc)
    camp = SurveyCampaign(semester=world["sem"], title="R", opens_at=now - timedelta(days=5),
                          closes_at=now + timedelta(days=2, hours=12))
    session.add(camp)
    await session.commit()
    assert await reminders.due(now) is None                                   # switched off
    await settings_store.save(session, reminders=True)
    await session.commit()
    settings_store.reset_cache()
    assert (await reminders.due(now))[1] == "3"
    await settings_store.save(session, reminders_sent={str(camp.id): ["3"]})
    await session.commit()
    settings_store.reset_cache()
    assert await reminders.due(now) is None
    assert (await reminders.due(now + timedelta(days=1, hours=18)))[1] == "1"  # 18 h before the end


# ---------------------------------------------------------------- Suhbatlar, Ko‘rinish, menu picture

async def test_conversations_page_is_for_the_programmer_and_replies_as_the_bot(boss, dev, session, monkeypatch):
    from app.models import BotChat, BotMessage
    from app.web.admin import chats
    session.add_all([BotUser(id=555, first_name="Aziz", username="aziz_k"),
                     BotChat(chat_id=555, name="Aziz Karimov", username="aziz_k", last_text="/start", last_dir="in")])
    await session.flush()
    session.add_all([
        BotMessage(chat_id=555, message_id=1, dir="in", kind="text", text="/start", reaction="🔥"),
        BotMessage(chat_id=555, message_id=2, dir="out", kind="text", text="<b>Salom</b> <script>alert(1)</script>",
                   buttons=[[{"t": "📅 Dars jadvali", "k": "cb"}]], origin="bot"),
        BotMessage(chat_id=555, message_id=2, dir="tap", kind="tap", text="📅 Dars jadvali"),
    ])
    await session.commit()
    assert (await boss.get("/admin/chats")).status_code == 404       # superadmins don't even see that it exists
    page = (await dev.get("/admin/chats?c=555")).text
    assert "Aziz Karimov" in page and "@aziz_k" in page and "📅 Dars jadvali" in page and "🔥" in page
    assert 'href="/admin/users/555"' in page                          # the account's page
    assert "<b>Salom</b>" in page and "<script>alert" not in page    # Telegram's formatting only
    assert (await dev.get("/admin/chats?q=Karimov")).text.count('class="chat-item ') == 1
    feed = (await dev.get("/admin/chats/555/feed", params={"v": "x"})).json()
    assert "Salom" in feed["html"]
    assert (await dev.get("/admin/chats/555/feed", params={"v": feed["v"]})).json() == {"v": feed["v"]}  # no change

    FakeBot.sent = []
    monkeypatch.setattr(chats.chatlog, "new_bot", FakeBot)
    r = (await dev.post("/admin/chats/555/send", data={"text": "Guruhingiz to‘g‘rimi?"})).json()
    assert r == {"ok": True} and FakeBot.sent == [(555, "text", "Guruhingiz to‘g‘rimi?")]
    assert (await dev.post("/admin/chats/555/send", data={"text": " "})).status_code == 400


async def test_display_settings_follow_the_account(boss, session):
    r = await boss.post("/admin/account/display", data={"text": "125", "theme": "dark", "contrast": "true",
                                                         "motion": "true", "spacing": "nope", "links": ""})
    assert r.json()["ok"]
    boss_admin = await session.scalar(select(Admin).where(Admin.username == "boss"))
    await session.refresh(boss_admin)
    assert boss_admin.prefs == {"text": 125, "theme": "dark", "contrast": True, "motion": True, "spacing": False,
                                "links": False, "hide_names": False, "group_numbers": False}
    page = (await boss.get("/admin/account")).text
    assert '<html lang="uz" data-text="125" data-theme="dark" data-contrast data-motion>' in page
    await boss.post("/admin/account/display", data={"text": "400", "theme": "pink"})  # not offered: back to normal
    await session.refresh(boss_admin)
    assert boss_admin.prefs["text"] == 100 and boss_admin.prefs["theme"] == "auto"
    assert '<html lang="uz" >' in (await boss.get("/admin/account")).text


async def test_menu_picture_and_channel_link(boss, session, monkeypatch, tmp_path):
    from app.web.admin import manage
    monkeypatch.setattr(manage, "MENU_PHOTO_FILE", tmp_path / "menu.jpg")
    assert (await boss.get("/admin/bot/menu-photo")).headers["content-type"] == "image/jpeg"   # the demo picture
    await boss.post("/admin/bot/menu-photo", files={"photo": ("my.png", png_bytes(), "image/png")})
    assert (tmp_path / "menu.jpg").read_bytes()[:2] == b"\xff\xd8"   # stored as a JPEG
    await boss.post("/admin/bot/menu-photo", files={"photo": ("bad.png", b"not a picture", "image/png")})
    assert (tmp_path / "menu.jpg").is_file()                         # a broken upload doesn't replace it
    await boss.post("/admin/bot/menu-photo/reset")
    assert not (tmp_path / "menu.jpg").exists()

    assert manage.channel_link("@iqtisodiyot_uz1") == "https://t.me/iqtisodiyot_uz1"
    assert manage.channel_link("t.me/iqtisodiyot_uz1") == "https://t.me/iqtisodiyot_uz1"
    assert manage.channel_link("") == "" and manage.channel_link("https://example.com/x") is None
    await boss.post("/admin/bot/settings", data={"channel_url": "@iqtisodiyot_uz1", "bot_reactions": "true"})
    settings_store.reset_cache()
    values = await settings_store.load(session, fresh=True)
    assert values["channel_url"] == "https://t.me/iqtisodiyot_uz1" and values["bot_reactions"]


async def test_linking_dashboard_accounts_to_telegram_and_the_admin_mini_app(boss, session, world):
    boss_admin = await session.scalar(select(Admin).where(Admin.username == "boss"))
    await settings_store.save(session, bot_username="talabaovozibot")
    await session.commit()
    settings_store.reset_cache()
    # One-time link
    await boss.post("/admin/account/telegram/link")
    await session.refresh(boss_admin)
    page = (await boss.get("/admin/account")).text
    assert f"https://t.me/talabaovozibot?start=adm_{boss_admin.tg_link_token}" in page
    # By username: cleaned, and two accounts can't share one
    await boss.post("/admin/account/telegram/username", data={"telegram_username": "@Dekan_Bey"})
    await session.refresh(boss_admin)
    assert boss_admin.telegram_username == "dekan_bey"
    viewer = await session.scalar(select(Admin).where(Admin.username == "viewer"))
    await boss.post(f"/admin/admins/{viewer.id}/edit", data={"telegram_username": "dekan_bey", "is_active": "true",
                                                              "role": "viewer"})
    await session.refresh(viewer)
    assert viewer.telegram_username is None
    await boss.post(f"/admin/admins/{viewer.id}/edit", data={"telegram_username": "t.me/kuzatuvchi1", "is_active": "true",
                                                              "role": "viewer"})
    await session.refresh(viewer)
    assert viewer.telegram_username == "kuzatuvchi1"

    # The Mini App: only for a linked, active account; Telegram's signature is checked.
    assert (await boss.get("/api/miniapp/admin")).status_code == 401
    tg = {"X-Telegram-Init-Data": sign_init_data({"id": 4242, "language_code": "ru"}, "123:TEST")}
    assert (await boss.get("/api/miniapp/admin", headers=tg)).status_code == 403
    viewer.telegram_id = 4242
    await session.commit()
    data = (await boss.get("/api/miniapp/admin", headers=tg)).json()
    assert data["lang"] == "ru" and data["admin"]["role"] and "stats" in data and data["links"]["admin"] == ""
    bad = {"X-Telegram-Init-Data": sign_init_data({"id": 4242}, "999:OTHER")}
    assert (await boss.get("/api/miniapp/admin", headers=bad)).status_code == 401
    assert (await boss.get("/miniapp/admin")).status_code == 200


async def test_hiding_who_rated_from_ones_own_panel(dev, session, world):
    """A programmer can hide who rated from their own panel, e.g. while showing the screen."""
    user, teacher, camp, fb = await rated_by_aziz(session, world, comment="Yaxshi o‘qituvchi")
    assert "Aziz Karimov" in (await dev.get("/admin/comments")).text
    await dev.post("/admin/account/display", data={"text": "100", "theme": "auto", "hide_names": "true"})
    page = (await dev.get("/admin/comments")).text
    assert "Yaxshi o‘qituvchi" in page and "Aziz Karimov" not in page
    assert "Aziz Karimov" not in (await dev.get(f"/admin/teachers/{teacher.id}")).text
    await dev.post("/admin/account/display", data={"text": "100", "theme": "auto"})
    assert "Aziz Karimov" in (await dev.get("/admin/comments")).text
