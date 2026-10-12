import re
import time
from datetime import date, datetime, timedelta, timezone

import pytest
from httpx import ASGITransport, AsyncClient

from app.faculty import is_faculty_group
from app.models import (Admin, BotUser, Feedback, Group, Semester, Subject, SurveyCampaign, Teacher,
                        TeachingAssignment)
from app.security import hash_password
from app.web.main import create_app
from app.web.telegram_auth import sign_init_data, verify_init_data

TOKEN = "123:TEST"
RATING = {"teaching": 8, "fairness": 7, "professionalism": 9}


def test_init_data_signature():
    good = sign_init_data({"id": 42, "first_name": "A"}, TOKEN)
    assert verify_init_data(good, TOKEN)["id"] == 42
    assert verify_init_data(good, "999:OTHER") is None                    # signed for another bot
    assert verify_init_data(good.replace("42", "43"), TOKEN) is None       # tampered
    old = sign_init_data({"id": 42}, TOKEN, auth_date=int(time.time()) - 2 * 86400)
    assert verify_init_data(old, TOKEN) is None                            # expired


def test_faculty_group_codes():
    assert is_faculty_group("25-16 IQT (ru)") and is_faculty_group("24-20 FTO‘IT") and is_faculty_group("26-12 FTO'IT")
    assert is_faculty_group("25-03 BB (ing)") and is_faculty_group("25-35a ATT")
    assert not is_faculty_group("25-24 B") and not is_faculty_group("23-12 SXMIM (ru)")


@pytest.fixture
async def world(session):
    now = datetime.now(timezone.utc)
    g = Group(name="25-16 IQT", course=2, student_count=20)
    sem = Semester(name="S", starts_on=date(2026, 9, 1), ends_on=date(2027, 1, 20), is_current=True)
    mine, other = Teacher(full_name="My Teacher"), Teacher(full_name="Not My Teacher")
    subj = Subject(name="Makro")
    camp = SurveyCampaign(semester=sem, title="Mid", opens_at=now - timedelta(days=1), closes_at=now + timedelta(days=1))
    session.add_all([g, camp, other, Admin(username="admin", password_hash=hash_password("secret123")),
                     TeachingAssignment(semester=sem, teacher=mine, group=g, subject=subj, lesson_type="lecture"),
                     TeachingAssignment(semester=sem, teacher=other, group=Group(name="25-17 IQT"), subject=subj)])
    await session.flush()
    user = BotUser(id=555, first_name="Aziz", username="aziz_k", group_id=g.id, onboarded_at=now)
    session.add(user)
    await session.commit()
    return {"user": user, "group": g, "mine": mine, "other": other, "campaign": camp}


@pytest.fixture
async def client(world):
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        yield c


def tg_headers(user_id=555, **extra):
    return {"X-Telegram-Init-Data": sign_init_data({"id": user_id, "first_name": "Aziz", **extra}, TOKEN)}


async def test_survey_requires_telegram_and_refuses_stopped_accounts(client, session, world):
    assert (await client.get("/api/survey")).status_code == 401
    assert (await client.get("/api/survey", headers={"X-Telegram-Init-Data": "hash=bad"})).status_code == 401
    # Someone who never started the bot gets an account (no group: every teacher) - with what Telegram told.
    r = await client.get("/api/survey", headers=tg_headers(999, photo_url="https://t.me/i/userpic/320/x.jpg"))
    assert r.status_code == 200 and r.json()["group"] is None
    assert {t["name"] for t in r.json()["teachers"]} == {"My Teacher", "Not My Teacher"}
    newcomer = await session.get(BotUser, 999)
    assert newcomer.first_name == "Aziz" and newcomer.profile["webapp_user"]["photo_url"].endswith("x.jpg")
    world["user"].blocked = True
    await session.commit()
    r = await client.get("/api/survey", headers=tg_headers())
    assert r.status_code == 403 and r.json()["detail"] == "blocked"
    world["user"].blocked = False
    world["user"].suspended_until = datetime.now(timezone.utc) + timedelta(hours=1)
    await session.commit()
    assert (await client.get("/api/survey", headers=tg_headers())).json()["detail"] == "suspended"


async def test_survey_flow(client, session, world):
    r = await client.get("/api/survey", headers=tg_headers())
    assert r.status_code == 200
    data = r.json()
    assert [t["name"] for t in data["teachers"]] == ["My Teacher"] and data["texts"]["title"]

    mine, other = world["mine"].id, world["other"].id
    assert (await client.post(f"/api/survey/teachers/{other}", json=RATING, headers=tg_headers())).status_code == 400
    assert (await client.post(f"/api/survey/teachers/{mine}", json={**RATING, "professionalism": 11},
                              headers=tg_headers())).status_code == 422
    assert (await client.post(f"/api/survey/teachers/{mine}", json={**RATING, "comment": " Zo'r! "},
                              headers=tg_headers())).status_code == 200
    # Submitting again updates the same evaluation instead of creating a duplicate.
    assert (await client.post(f"/api/survey/teachers/{mine}", json={**RATING, "teaching": 10},
                              headers=tg_headers())).status_code == 200
    from sqlalchemy import select
    rows = (await session.execute(select(Feedback))).scalars().all()
    assert len(rows) == 1 and rows[0].teaching == 10 and rows[0].user_id == 555
    assert rows[0].group_id == world["group"].id

    r = await client.get("/api/survey", headers={"X-Dev-User": "555"})  # dev mode
    assert r.json()["teachers"][0]["feedback"]["teaching"] == 10 and r.json()["group"] == "25-16 IQT"

    # Without a group: any teacher of this semester, rated as "no group".
    world["user"].group_id = None
    await session.commit()
    assert (await client.post(f"/api/survey/teachers/{other}", json=RATING, headers=tg_headers())).status_code == 200
    fb = await session.scalar(select(Feedback).where(Feedback.teacher_id == other))
    assert fb.group_id is None


async def test_too_many_submits_are_refused(client, world, monkeypatch):
    from app.web import survey
    monkeypatch.setattr(survey, "SUBMITS_PER_MINUTE", 3)
    survey._submits.clear()
    codes = [(await client.post(f"/api/survey/teachers/{world['mine'].id}", json=RATING, headers=tg_headers())).status_code
             for _ in range(5)]
    assert codes == [200, 200, 200, 429, 429]


async def test_closed_campaign_rejects(client, session, world):
    world["campaign"].closes_at = datetime.now(timezone.utc) - timedelta(minutes=1)
    await session.commit()
    r = await client.post(f"/api/survey/teachers/{world['mine'].id}", json=RATING, headers=tg_headers())
    assert r.status_code == 400 and r.json()["detail"] == "no_campaign"


async def test_dashboard_pages(client, world):
    assert (await client.get("/admin")).status_code == 303  # must log in
    r = await client.post("/admin/login", data={"username": "admin", "password": "wrong"})
    assert r.status_code == 401
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    await client.post(f"/api/survey/teachers/{world['mine'].id}", json={**RATING, "comment": "Great"},
                      headers=tg_headers())
    for url in ["/admin", "/admin?campaign=all", "/admin/teachers", f"/admin/teachers/{world['mine'].id}",
                "/admin/groups", "/admin/comments", "/admin/comments?band=hi&q=Great", "/admin/export/comments.xlsx"]:
        r = await client.get(url)
        assert r.status_code == 200, url
    # The student list is gone; old links lead to the groups. Who rated is not on any page (not a programmer).
    assert (await client.get("/admin/students", follow_redirects=False)).headers["location"] == "/admin/groups"
    for url in ["/admin", "/admin/comments", f"/admin/teachers/{world['mine'].id}"]:
        html = (await client.get(url)).text
        assert "Aziz" not in html and "aziz_k" not in html and "555" not in html, url
    assert (await client.get("/admin/users")).status_code == 404               # programmer accounts only


async def test_blank_filters_mean_all(client, world):
    """Choosing "All groups" etc. submits an empty value (?group=); that must show everything, not an error."""
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    for url in ["/admin/comments?q=&teacher=&group=&page=", "/admin/teachers?department=", "/admin/groups?course=",
                "/admin/export/comments.xlsx?teacher=&group="]:
        assert (await client.get(url)).status_code == 200, url


async def test_overview_tiles_and_ranking(client, world):
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    await client.post(f"/api/survey/teachers/{world['mine'].id}", json={**RATING, "comment": "Great"},
                      headers=tg_headers())
    html = (await client.get("/admin")).text
    assert "1 / 1 ta bahoda izoh bor" in html                       # exact numbers behind the comments ring
    assert "Ishtirok voronkasi" not in html and "Talabalar ishtiroki" not in html
    assert 'class="rank-row"' in html and "My Teacher" in html         # most rated: 1 evaluation, rated 8.00
    assert "Admin sayti" in html


async def test_breadcrumbs_fallback(client, world):
    """Opened directly, a detail page shows its usual place: O‘qituvchilar › name (crumbs.js refines it)."""
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    html = (await client.get(f"/admin/teachers/{world['mine'].id}")).text
    assert 'class="head crumbed"' in html and '<span aria-current="page">My Teacher</span>' in html
    assert '"entity": "teacher:%d"' % world["mine"].id in html


async def test_comments_excel_follows_filters(client, world):
    from io import BytesIO
    from openpyxl import load_workbook
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    await client.post(f"/api/survey/teachers/{world['mine'].id}", json={**RATING, "comment": "Great"},
                      headers=tg_headers())
    rows = lambda r: list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))  # noqa: E731
    got = rows(await client.get("/admin/export/comments.xlsx"))
    assert len(got) == 2 and got[1][-1] == "Great" and got[1][5:8] == (8, 7, 9)       # no "who rated" columns
    assert "Aziz" not in str(got)
    assert len(rows(await client.get("/admin/export/comments.xlsx?band=lo"))) == 1   # header only
    assert len(rows(await client.get("/admin/export/comments.xlsx?q=nothing"))) == 1
    for gone in ["/admin/export/feedback.xlsx", "/admin/export/teachers.xlsx", "/admin/export/groups.xlsx"]:
        assert (await client.get(gone)).status_code == 404


async def test_comment_filters_narrow_each_other(client, world):
    """With a group chosen only its teachers are offered, with a teacher only their groups."""
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    group_id = world["group"].id
    html = (await client.get(f"/admin/comments?group={group_id}")).text
    assert ">My Teacher</option>" in html and "Not My Teacher</option>" not in html
    html = (await client.get(f"/admin/comments?teacher={world['other'].id}")).text
    assert ">25-17 IQT</option>" in html and ">25-16 IQT</option>" not in html
    html = (await client.get("/admin/comments")).text
    assert ">25-17 IQT</option>" in html and ">25-16 IQT</option>" in html and "Not My Teacher</option>" in html


async def test_groups_page_and_menu(client, world):
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    html = (await client.get("/admin/groups?course=3")).text
    assert "25-16 IQT" in html                                         # the course filter is the table's own (browser)
    assert 'class="row-link" href="/admin/comments?' in html           # the group name opens its comments
    assert "/admin/students" not in html and "/admin/users" not in html  # no student list, no Users item (viewer)
    assert re.search(r'href="/admin/groups\?[^"]*" class="on" data-nav>', html)


def test_specialty_names():
    from app.faculty import specialty_name
    assert specialty_name("FTO‘IT") == specialty_name("FTO'IT") == "Filologiya va tillarni o‘qitish: ingliz tili"
    assert specialty_name("BXA", "en") == "Accounting" and specialty_name("TUR", "ru") == "Туризм и гостеприимство"
    assert specialty_name("XYZ") == "XYZ"


async def test_dashboard_hidden_from_public_hostname(world):
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="https://abc.trycloudflare.com") as c:
        assert (await c.get("/admin/login")).status_code == 404
        assert (await c.get("/survey")).status_code == 200


async def test_dev_shortcut_and_dashboard_blocked_through_tunnel(world):
    """cloudflared forwards requests to 127.0.0.1 but adds Cf-Connecting-Ip - those must be treated as public."""
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test") as c:
        tunnel = {"Cf-Connecting-Ip": "1.2.3.4"}
        assert (await c.get("/api/survey", headers={"X-Dev-User": "555", **tunnel})).status_code == 401
        assert (await c.get("/admin/login", headers=tunnel)).status_code == 404
        assert (await c.get("/api/survey", headers={"X-Dev-User": "555"})).status_code == 200


async def test_dashboard_languages(client, world):
    await client.post("/admin/login", data={"username": "admin", "password": "secret123"})
    assert "Umumiy ko‘rinish" in (await client.get("/admin")).text          # Uzbek by default
    r = await client.get("/admin/lang/ru?next=/admin/teachers", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/admin/teachers"
    client.cookies.set("admin_lang", "ru")
    assert "Преподаватели" in (await client.get("/admin/teachers")).text
    client.cookies.set("admin_lang", "en")
    assert "Sign in" in (await client.get("/admin/login")).text
    r = await client.get("/admin/lang/en?next=https://evil.example", follow_redirects=False)
    assert r.headers["location"] == "/admin"                                    # no open redirect


def test_admin_translations_complete():
    from app.web.admin.i18n import TEXTS
    import re
    keys = set(TEXTS["uz"])
    for lang, texts in TEXTS.items():
        assert set(texts) == keys, (lang, keys ^ set(texts))
        for k, v in texts.items():
            assert set(re.findall(r"{(\w+)}", v)) == set(re.findall(r"{(\w+)}", TEXTS["uz"][k])), (lang, k)
